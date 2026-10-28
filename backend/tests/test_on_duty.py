"""On duty (plan 13): client locations, on-duty requests, client-site geofence for attendance, calendar, export.
2026-10-21 is a Wednesday; office timezone Asia/Kolkata = UTC+5:30. The office is at 12.9716, 77.5946 (default settings).
SQLite: the EXCLUDE constraint is not exercised here (the service pre-checks are)."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import get_settings
from app.models import Attendance, AuditLog, ClientLocation, LeaveType, OnDutyRequest, Setting, User
from app.services import attendance as att_svc
from app.services.notifications import OUTBOX
from tests.conftest import USERS, bearer

OFFICE = {"lat": 12.9716, "lng": 77.5946, "accuracy_m": 10}
SITE = {"lat": 13.0827, "lng": 80.2707}  # client site, far from the office
AT_SITE = {"lat": 13.0827 + 0.0009, "lng": 80.2707, "accuracy_m": 10}  # ~100 m north of the site
NEAR_SITE_OUT = {"lat": 13.0827 + 0.0030, "lng": 80.2707, "accuracy_m": 10}  # ~334 m north: outside a 200 m radius
NOW = dt.datetime(2026, 10, 21, 3, 30, tzinfo=dt.UTC)  # 09:00 IST


def _set_clock(monkeypatch, when: dt.datetime) -> None:
    for mod in ("leave", "attendance", "late"):
        monkeypatch.setattr(f"app.services.{mod}.utcnow", lambda w=when: w)


@pytest.fixture(autouse=True)
def clock(monkeypatch):
    _set_clock(monkeypatch, NOW)
    OUTBOX.clear()
    st = get_settings()
    monkeypatch.setattr(st, "leave_write_rate_per_min", 1000)
    monkeypatch.setattr(st, "leave_write_rate_per_day", 100000)
    monkeypatch.setattr(st, "leave_decide_rate_per_min", 1000)
    return lambda when: _set_clock(monkeypatch, when)


@pytest.fixture
async def H(tokens_for):
    cache: dict = {}

    async def _h(role):
        if role not in cache:
            cache[role] = bearer(await tokens_for(role))
        return cache[role]

    return _h


@pytest.fixture
async def site(client, H):
    r = await client.post(
        "/api/admin/client-locations",
        headers=await H("admin"),
        json={"name": "Acme Chennai", "address": "Mount Road, Chennai", **SITE, "radius_m": 200},
    )
    assert r.status_code == 201, r.text
    return r.json()


def od(site_id, f="2026-10-21", t=None, reason="client installation"):
    return {"client_location_id": site_id, "from_date": f, "to_date": t or f, "reason": reason}


async def apply(client, h, site_id, f="2026-10-21", t=None):
    return await client.post("/api/on-duty/requests", headers=h, json=od(site_id, f, t))


async def approved(client, H, site_id, f="2026-10-21", t=None) -> dict:
    r = await apply(client, await H("employee"), site_id, f, t)
    assert r.status_code == 201, r.text
    a = await client.put(f"/api/admin/requests/on_duty/{r.json()['id']}/approve", headers=await H("manager"), json={})
    assert a.status_code == 200, a.text
    return a.json()


def code(r):
    return r.json()["error"]["code"]


def fix():
    return {"fix_time": att_svc.utcnow().isoformat()}


async def check_in(client, h, where):
    return await client.post("/api/attendance/check-in", headers=h, json={**where, **fix()})


async def check_out(client, h, where):
    return await client.post("/api/attendance/check-out", headers=h, json={**where, **fix()})


# ------------------------------------------------------------------ client locations (admin)
async def test_client_location_crud_and_audit(client, H, site, db):
    h = await H("hr")  # hr holds client_location.manage too
    assert site["lat"] == SITE["lat"] and site["radius_m"] == 200 and site["is_active"] is True
    lst = (await client.get("/api/admin/client-locations", headers=h)).json()
    assert [x["name"] for x in lst["items"]] == ["Acme Chennai"]
    upd = {"name": "Acme Chennai HQ", "address": None, **SITE, "radius_m": 300, "is_active": True}
    r = await client.put(f"/api/admin/client-locations/{site['id']}", headers=h, json=upd)
    assert r.status_code == 200 and r.json()["radius_m"] == 300 and r.json()["address"] is None
    dup = await client.post("/api/admin/client-locations", headers=h, json={**upd, "name": "acme chennai hq"})
    assert dup.status_code == 409
    actions = {a for a in (await db.execute(select(AuditLog.action))).scalars()}
    assert {"client_location.create", "client_location.update"} <= actions
    r = await client.delete(f"/api/admin/client-locations/{site['id']}", headers=h)
    assert r.status_code == 204
    assert (await client.get(f"/api/admin/client-locations/{site['id']}", headers=h)).status_code == 404


@pytest.mark.parametrize(
    "bad",
    [
        {"lat": 91},
        {"lng": -181},
        {"radius_m": 10},
        {"radius_m": 6000},
        {"lat": 0, "lng": 0},
        {"name": "  "},
        {"extra": 1},
    ],
)
async def test_client_location_validation(client, H, bad):
    body = {"name": "Site", **SITE, "radius_m": 200, **bad}
    r = await client.post("/api/admin/client-locations", headers=await H("admin"), json=body)
    assert r.status_code == 422, r.text


@pytest.mark.parametrize("role", ["employee", "manager"])
async def test_client_location_admin_endpoints_need_permission(client, H, site, role):
    h = await H(role)
    assert (await client.get("/api/admin/client-locations", headers=h)).status_code == 403
    assert (await client.get(f"/api/admin/client-locations/{site['id']}", headers=h)).status_code == 403
    r = await client.post("/api/admin/client-locations", headers=h, json={"name": "X", **SITE})
    assert r.status_code == 403


async def test_employee_picker_has_no_coordinates_and_only_active(client, H, site):
    h = await H("admin")
    other = {"name": "Old Site", **SITE, "radius_m": 200, "is_active": False}
    assert (await client.post("/api/admin/client-locations", headers=h, json=other)).status_code == 201
    items = (await client.get("/api/on-duty/locations", headers=await H("employee"))).json()["items"]
    assert items == [{"id": site["id"], "name": "Acme Chennai", "address": "Mount Road, Chennai"}]


async def test_location_in_use_cannot_be_deleted(client, H, site):
    assert (await apply(client, await H("employee"), site["id"])).status_code == 201
    r = await client.delete(f"/api/admin/client-locations/{site['id']}", headers=await H("admin"))
    assert r.status_code == 409 and code(r) == "IN_USE"


# ------------------------------------------------------------------ apply / approve
async def test_apply_list_and_approve(client, H, site, db):
    h = await H("employee")
    r = await apply(client, h, site["id"], "2026-10-21", "2026-10-26")  # Wed..Mon: Sat/Sun are weekly offs
    assert r.status_code == 201, r.text
    b = r.json()
    assert b["status"] == "pending" and b["days"] == 4
    assert b["client_location"] == {"id": site["id"], "name": "Acme Chennai", "address": "Mount Road, Chennai"}
    assert "lat" not in b["client_location"]
    assert OUTBOX and OUTBOX[-1]["payload"]["title"] == "New on duty request"
    mine = (await client.get("/api/on-duty/requests", headers=h)).json()
    assert [x["id"] for x in mine["items"]] == [b["id"]]
    assert (await client.get(f"/api/on-duty/requests/{b['id']}", headers=h)).json()["reason"] == "client installation"
    team = (await client.get("/api/admin/requests", headers=await H("manager"), params={"type": "on_duty"})).json()
    assert team["items"][0]["type"] == "on_duty" and team["items"][0]["user"]["employee_id"] == USERS["employee"]
    a = await client.put(f"/api/admin/requests/on_duty/{b['id']}/approve", headers=await H("manager"), json={"comment": "ok"})
    assert a.status_code == 200 and a.json()["status"] == "approved"
    assert OUTBOX[-1]["payload"]["title"] == "On duty request approved"
    row = await db.get(OnDutyRequest, __import__("uuid").UUID(b["id"]))
    assert row.reason == "client installation"  # decrypted transparently (stored as reason_enc)


async def test_reject_and_self_approval_blocked(client, H, site):
    r = await apply(client, await H("manager"), site["id"])
    assert r.status_code == 201
    own = await client.put(f"/api/admin/requests/on_duty/{r.json()['id']}/approve", headers=await H("manager"), json={})
    assert own.status_code == 403
    rej = await client.put(
        f"/api/admin/requests/on_duty/{r.json()['id']}/reject", headers=await H("hr"), json={"comment": "not needed"}
    )
    assert rej.status_code == 200 and rej.json()["status"] == "rejected"


@pytest.mark.parametrize(
    ("f", "t", "err"),
    [
        ("2026-10-20", None, "INVALID_DATE_RANGE"),  # yesterday
        ("2026-10-24", "2026-10-25", "ZERO_WORKING_DAYS"),  # weekend only: no working days
        ("2026-10-21", "2026-11-30", "INVALID_DATE_RANGE"),  # more than 31 days
    ],
)
async def test_apply_date_rules(client, H, site, f, t, err):
    r = await apply(client, await H("employee"), site["id"], f, t)
    assert r.status_code in (400, 409, 422) and code(r) == err, r.text


async def test_to_before_from_is_422(client, H, site):
    r = await apply(client, await H("employee"), site["id"], "2026-10-22", "2026-10-21")
    assert r.status_code == 422


async def test_inactive_or_unknown_location_refused(client, H, site):
    h = await H("admin")
    body = {"name": site["name"], **SITE, "radius_m": 200, "is_active": False}
    assert (await client.put(f"/api/admin/client-locations/{site['id']}", headers=h, json=body)).status_code == 200
    e = await H("employee")
    assert code(await apply(client, e, site["id"])) == "VALIDATION_ERROR"
    assert code(await apply(client, e, "00000000-0000-0000-0000-000000000000")) == "VALIDATION_ERROR"


async def test_overlaps_with_on_duty_and_leave_both_ways(client, H, site, db):
    lt = LeaveType(code="CL", name="Casual Leave", annual_quota=12)
    db.add(lt)
    await db.commit()
    h = await H("employee")
    assert (await apply(client, h, site["id"], "2026-10-22")).status_code == 201
    assert code(await apply(client, h, site["id"], "2026-10-21", "2026-10-23")) == "OVERLAPPING_REQUEST"
    leave = {"leave_type_id": str(lt.id), "from_date": "2026-10-22", "to_date": "2026-10-22", "reason": "x"}
    r = await client.post("/api/leave/requests", headers=h, json=leave)
    assert r.status_code == 409 and code(r) == "OVERLAPPING_REQUEST" and "on-duty" in r.json()["error"]["message"]
    leave["from_date"] = leave["to_date"] = "2026-10-27"
    assert (await client.post("/api/leave/requests", headers=h, json=leave)).status_code == 201
    assert code(await apply(client, h, site["id"], "2026-10-27")) == "OVERLAPPING_REQUEST"


async def test_cancel_pending_then_reapply(client, H, site):
    h = await H("employee")
    r = await apply(client, h, site["id"])
    c = await client.post(f"/api/on-duty/requests/{r.json()['id']}/cancel", headers=h)
    assert c.status_code == 200 and c.json()["status"] == "cancelled"
    assert (await apply(client, h, site["id"])).status_code == 201


async def test_other_employee_cannot_see_or_cancel(client, H, site):
    r = await apply(client, await H("employee"), site["id"])
    other = await H("hr")
    assert (await client.get(f"/api/on-duty/requests/{r.json()['id']}", headers=other)).status_code == 404
    assert (await client.post(f"/api/on-duty/requests/{r.json()['id']}/cancel", headers=other)).status_code == 404


async def test_deactivation_cancels_pending_on_duty(client, H, site, db):
    r = await apply(client, await H("employee"), site["id"])
    uid = (await db.execute(select(User.id).where(User.employee_id == USERS["employee"]))).scalar_one()
    d = await client.post(f"/api/admin/employees/{uid}/deactivate", headers=await H("admin"), json={})
    assert d.status_code == 200, d.text
    row = await db.get(OnDutyRequest, __import__("uuid").UUID(r.json()["id"]))
    await db.refresh(row)
    assert row.status == "cancelled"


# ------------------------------------------------------------------ attendance at the client site
async def test_approved_on_duty_requires_client_site_for_check_in(client, H, site, db):
    await approved(client, H, site["id"])
    h = await H("employee")
    today = (await client.get("/api/attendance/today", headers=h)).json()
    assert today["record"] is None
    assert today["on_duty"] == {
        "id": site["id"], "name": "Acme Chennai", "address": "Mount Road, Chennai", **SITE, "radius_m": 200,
    }  # fmt: skip
    r = await check_in(client, h, OFFICE)  # the office does not count on an on-duty day
    assert r.status_code == 422 and code(r) == "OUTSIDE_CLIENT_LOCATION", r.text
    assert "Acme Chennai" in r.json()["error"]["message"]
    r = await check_in(client, h, NEAR_SITE_OUT)
    assert code(r) == "OUTSIDE_CLIENT_LOCATION"
    r = await check_in(client, h, AT_SITE)
    assert r.status_code == 201, r.text
    b = r.json()
    assert b["within_fence"] is True and b["client_location"] == {"id": site["id"], "name": "Acme Chennai"}
    rec = (await db.execute(select(Attendance))).scalar_one()
    assert str(rec.client_location_id) == site["id"] and rec.check_in_distance_m < 200
    rejected = (await db.execute(select(AuditLog).where(AuditLog.action == "attendance.punch_rejected"))).scalars().all()
    assert [a.after["reason"] for a in rejected] == ["OUTSIDE_CLIENT_LOCATION", "OUTSIDE_CLIENT_LOCATION"]
    t = (await client.get("/api/attendance/today", headers=h)).json()
    assert t["client_location"]["name"] == "Acme Chennai" and t["on_duty"]["id"] == site["id"]


async def test_pending_or_other_day_on_duty_keeps_office_fence(client, H, site):
    h = await H("employee")
    assert (await apply(client, h, site["id"])).status_code == 201  # pending only
    await approved(client, H, site["id"], "2026-10-22")  # tomorrow
    r = await check_in(client, h, OFFICE)
    assert r.status_code == 201 and r.json()["client_location"] is None
    assert (await client.get("/api/attendance/today", headers=h)).json()["on_duty"] is None


async def test_checkout_geofence_uses_client_site(client, H, site, db, clock):
    db.add(Setting(key="attendance.checkout_geofence", value=True))
    await db.commit()
    await approved(client, H, site["id"])
    h = await H("employee")
    assert (await check_in(client, h, AT_SITE)).status_code == 201
    clock(NOW + dt.timedelta(hours=9))
    r = await check_out(client, h, OFFICE)
    assert r.status_code == 422 and code(r) == "OUTSIDE_CLIENT_LOCATION"
    r = await check_out(client, h, AT_SITE)
    assert r.status_code == 200, r.text
    assert r.json()["check_out_at"] is not None


async def test_late_rules_still_apply_on_duty(client, H, site, clock):
    await approved(client, H, site["id"])
    clock(dt.datetime(2026, 10, 21, 4, 30, tzinfo=dt.UTC))  # 10:00 IST, shift 09:30 + 10 min grace
    r = await check_in(client, await H("employee"), AT_SITE)
    assert r.status_code == 201 and r.json()["status"] == "late" and r.json()["minutes_late"] == 30


async def test_calendar_marks_on_duty_days(client, H, site):
    await approved(client, H, site["id"], "2026-10-21", "2026-10-22")
    days = (await client.get("/api/attendance/calendar", headers=await H("employee"), params={"month": "2026-10"})).json()
    by = {d["date"]: d for d in days["days"]}
    assert by["2026-10-21"]["on_duty"] == "Acme Chennai" and by["2026-10-22"]["on_duty"] == "Acme Chennai"
    assert by["2026-10-23"]["on_duty"] is None


async def test_leave_approval_blocked_by_approved_on_duty(client, H, site, db):
    lt = LeaveType(code="CL", name="Casual Leave", annual_quota=12)
    db.add(lt)
    await db.commit()
    h = await H("employee")
    leave = {"leave_type_id": str(lt.id), "from_date": "2026-10-22", "to_date": "2026-10-22", "reason": "x"}
    lr = await client.post("/api/leave/requests", headers=h, json=leave)
    assert lr.status_code == 201
    # an on-duty request on the same day can only exist if it was created before the leave (e.g. by data import)
    uid = (await db.execute(select(User.id).where(User.employee_id == USERS["employee"]))).scalar_one()
    loc = (await db.execute(select(ClientLocation))).scalar_one()
    db.add(OnDutyRequest(user_id=uid, client_location_id=loc.id, from_date=dt.date(2026, 10, 22),
                         to_date=dt.date(2026, 10, 22), days=1, reason="x", status="approved"))  # fmt: skip
    await db.commit()
    a = await client.put(f"/api/admin/requests/leave/{lr.json()['id']}/approve", headers=await H("manager"), json={})
    assert a.status_code == 409 and code(a) == "OVERLAPPING_REQUEST"


# ------------------------------------------------------------------ export
async def test_export_on_duty_sheet_and_attendance_column(client, H, site):
    await approved(client, H, site["id"])
    assert (await check_in(client, await H("employee"), AT_SITE)).status_code == 201
    r = await client.get(
        "/api/admin/export", headers=await H("admin"), params={"type": "all", "from": "2026-10-21", "to": "2026-10-21"}
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert "On Duty" in wb.sheetnames
    ws = wb["On Duty"]
    head = [c.value for c in ws[1]]
    assert "Reason" not in head  # sensitive columns only with include_reasons
    row = [c.value for c in next(ws.iter_rows(min_row=2))]
    assert row[head.index("Client Location")] == "Acme Chennai" and row[head.index("Status")] == "approved"
    att = wb["Attendance"]
    ahead = [c.value for c in att[1]]
    arow = [c.value for c in next(att.iter_rows(min_row=2))]
    assert arow[ahead.index("Client Location")] == "Acme Chennai"
    one = await client.get(
        "/api/admin/export", headers=await H("hr"), params={"type": "on_duty", "from": "2026-10-21", "to": "2026-10-21"}
    )
    assert load_workbook(io.BytesIO(one.content)).sheetnames == ["On Duty"]


async def test_retention_purges_decided_on_duty_text(client, H, site, db, clock):
    from app.services import leave as lv

    r = await approved(client, H, site["id"])
    db.add(Setting(key="leave.reason_retention_days", value=30))
    await db.commit()
    assert (await lv.purge_old_text(db))["on_duty"] == 0  # too recent
    clock(NOW + dt.timedelta(days=40))
    assert (await lv.purge_old_text(db))["on_duty"] == 1
    db.expire_all()
    row = await db.get(OnDutyRequest, __import__("uuid").UUID(r["id"]))
    assert row.reason is None and row.decision_comment is None and row.status == "approved"
