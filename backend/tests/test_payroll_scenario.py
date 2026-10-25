"""Payroll: a realistic month for six people, end to end (QA scenario).

June 2025 (30 calendar days, PT is half-yearly so none is due in June). Expected amounts were worked out by hand from the
Indian rules, NOT by calling the application's calculator:

* PF = 12 % of basic, basic capped at 15,000 (so 1,800 at most) unless the structure says "PF on full basic".
* ESI = 0.75 % (employee) / 3.25 % (employer) of ESI wages when the monthly gross is <= 20,000, rounded UP to the next rupee.
* Business rule: the employer PF and ESI share is ALSO deducted from the employee (payroll.employer_share_from_employee, default
  on), as lines PF_ER_REC / ESI_ER_REC equal to the employer contribution; net pay falls by both shares.
* Every other computed line is rounded half-up to the rupee.

    EMP001  basic 20,000 + HRA 8,000, TDS 1,200, LOP override 1.5 days, arrears 2,500
            paid 28.5/30 = 0.95  ->  basic 19,000, HRA 7,600, +2,500 = gross 29,100
            PF 12 % x min(19,000, 15,000) = 1,800; no ESI (gross 28,000 > 20,000); employer PF share 1,800; TDS 1,200  -> net 29,100 - 4,800 = 24,300
    MGR001  basic 30,000 + HRA 10,000, PF on FULL basic (3,600), TDS 3,000, canteen 500
            employer PF share 3,600; gross 40,000, deductions 10,700  -> net 29,300
    LOW001  basic 12,000 + HRA 4,800 + conveyance 2,050 = 18,850 (<= 20,000, ESI applies)
            PF 1,440; ESI 18,850 x 0.75 % = 141.375 -> 142 (round UP); employer ESI 612.6 -> 613; employer PF share 1,440
            deductions 1,440 + 1,440 + 142 + 613 = 3,635  -> net 15,215
    JOIN01  joined 11 June: 20/30 days. basic 15,000 + HRA 5,000 (= 20,000 exactly, ESI applies at the threshold)
            basic 10,000, HRA 3,333.33 -> 3,333, gross 13,333; PF 1,200; ESI 13,333 x 0.75 % = 99.9975 -> 100 (round UP); employer ESI 433.32 -> 434
            deductions 1,200 + 1,200 + 100 + 434 = 2,934  -> net 10,399
    LEAV01  last day 15 June: 15/30 days. basic 40,000 + HRA 10,000, TDS 2,000, advance recovery 1,000
            basic 20,000, HRA 5,000, gross 25,000; PF 1,800 (capped); no ESI; deductions 1,800 + 1,800 + 2,000 + 1,000 = 6,600 -> net 18,400
    HR001   basic 50,000, nothing else: PF 1,800 + employer share 1,800, net 46,400

    run totals: gross 176,283, deductions 32,269 (19,582 + 12,687 employer share recovered), net 144,014, employer PF 11,640, employer ESI 1,047,
    company cost for the employer shares 0 (remitted 12,687, recovered 12,687)
"""

import datetime as dt
import io
import uuid
from decimal import ROUND_CEILING, ROUND_HALF_UP
from decimal import Decimal as D

import pytest_asyncio
from openpyxl import load_workbook

from app.models import EmployeeProfile, User
from tests.conftest import PASSWORD
from tests.payroll_fixtures import (
    PERIOD,
    add_structure,
    exclude_from_payroll,
    new_run,
    publish_run,
    row_for,
    set_setting,
)

PW = PASSWORD


# ------------------------------------------------------------------------------------------------ independent oracle
def oracle_pf(basic: D, full_basic: bool = False) -> D:
    wage = basic if full_basic else min(basic, D(15000))
    return (wage * D("0.12")).quantize(D(1), ROUND_HALF_UP)


def oracle_esi(wage: D, rate: str) -> D:
    return (wage * D(rate) / 100).quantize(D(1), ROUND_CEILING)


def test_oracle_matches_the_hand_worked_numbers():
    assert oracle_pf(D(19000)) == 1800 and oracle_pf(D(30000), True) == 3600 and oracle_pf(D(12000)) == 1440
    assert oracle_esi(D(18850), "0.75") == 142 and oracle_esi(D(18850), "3.25") == 613
    assert oracle_esi(D(13333), "0.75") == 100 and oracle_esi(D(13333), "3.25") == 434


# ------------------------------------------------------------------------------------------------ world
async def _make_user(db, seeded, eid, name, join=None):
    u = User(
        employee_id=eid,
        password_hash="x",
        full_name=name,
        official_email=f"{eid.lower()}@corp-mail.com",
        role_id=seeded["roles"]["employee"],
        department_id=seeded["dept"],
        shift_id=seeded["shift"],
        must_change_password=False,
        status="active",
        join_date=join,
    )
    db.add(u)
    await db.flush()
    db.add(EmployeeProfile(user_id=u.id))
    await db.commit()
    return str(u.id)


@pytest_asyncio.fixture
async def world(client, H, uid, seeded, db):
    """Six paid people (admin is the payroll operator and is not paid). Returns {employee_id: user_id}."""
    await set_setting(client, H["admin"], lop_absent_days=False)
    await exclude_from_payroll(client, H["hr"], uid["admin"])
    ids = {"EMP001": uid["employee"], "MGR001": uid["manager"], "HR001": uid["hr"]}
    ids["LOW001"] = await _make_user(db, seeded, "LOW001", "Low Paid")
    ids["JOIN01"] = await _make_user(db, seeded, "JOIN01", "New Joiner", join=dt.date(2025, 6, 11))
    ids["LEAV01"] = await _make_user(db, seeded, "LEAV01", "Leaver")
    leaver = await db.get(User, uuid.UUID(ids["LEAV01"]))
    leaver.status, leaver.deactivated_at, leaver.deactivation_reason = (
        "inactive",
        dt.datetime(2025, 6, 16, 5, 0, tzinfo=dt.UTC),
        "resigned",
    )
    await db.commit()
    a = H["admin"]
    await add_structure(client, a, ids["EMP001"], basic="20000", hra="8000", tds_monthly="1200")
    await add_structure(client, a, ids["MGR001"], basic="30000", hra="10000", tds_monthly="3000", pf_on_full_basic=True,
                        fixed_deductions=[{"code": "CANTEEN", "name": "Canteen", "amount": "500"}])  # fmt: skip
    await add_structure(client, a, ids["LOW001"], basic="12000", hra="4800",
                        allowances=[{"code": "CONV", "name": "Conveyance", "amount": "2050", "prorate": True, "esi_wage": True}])  # fmt: skip
    await add_structure(client, a, ids["JOIN01"], basic="15000", hra="5000")
    await add_structure(client, a, ids["LEAV01"], basic="40000", hra="10000", tds_monthly="2000")
    await add_structure(client, a, ids["HR001"], basic="50000")
    return ids


def _lines(p, key):
    return {x["code"]: D(x["amount"]) for x in p[key]}


async def _admin_payslip(client, hdr, run_id, eid):
    row = await row_for(client, hdr, run_id, eid)
    r = await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=hdr)
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------------------------------------------ the month
async def test_full_month_run_for_six_people(client, H, world):
    admin = H["admin"]
    run = await new_run(client, admin)
    assert run["counts"]["employees"] == 6
    # adjustments / overrides entered by the operator
    emp = await _admin_payslip(client, admin, run["id"], "EMP001")
    r = await client.put(
        f"/api/admin/payroll/payslips/{emp['id']}",
        json={"expected_version": emp["version"], "lop_override": 1.5, "lop_override_reason": "Unpaid leave 9-10 June, 11th half day",
              "adjustments": [{"code": "ARREARS", "label": "May increment arrears", "amount": "2500"}]},
        headers=admin,
    )  # fmt: skip
    assert r.status_code == 200, r.text
    leaver = await _admin_payslip(client, admin, run["id"], "LEAV01")
    r = await client.put(
        f"/api/admin/payroll/payslips/{leaver['id']}",
        json={"expected_version": leaver["version"], "adjustments": [{"code": "ADVANCE_RECOVERY", "label": "Salary advance", "amount": "1000"}]},
        headers=admin,
    )  # fmt: skip
    assert r.status_code == 200, r.text

    expect = {  # eid -> (paid_days, gross, {deduction code: amount}, net); *_ER_REC = employer share recovered
        "EMP001": (28.5, "29100.00", {"PF": 1800, "PF_ER_REC": 1800, "TDS": 1200}, "24300.00"),
        "MGR001": (30.0, "40000.00", {"PF": 3600, "PF_ER_REC": 3600, "TDS": 3000, "CANTEEN": 500}, "29300.00"),
        "LOW001": (30.0, "18850.00", {"PF": 1440, "PF_ER_REC": 1440, "ESI": 142, "ESI_ER_REC": 613}, "15215.00"),
        "JOIN01": (20.0, "13333.00", {"PF": 1200, "PF_ER_REC": 1200, "ESI": 100, "ESI_ER_REC": 434}, "10399.00"),
        "LEAV01": (15.0, "25000.00", {"PF": 1800, "PF_ER_REC": 1800, "TDS": 2000, "ADVANCE_RECOVERY": 1000}, "18400.00"),
        "HR001": (30.0, "50000.00", {"PF": 1800, "PF_ER_REC": 1800}, "46400.00"),
    }
    for eid, (paid, gross, ded, net) in expect.items():
        p = await _admin_payslip(client, admin, run["id"], eid)
        assert p["days"]["paid_days"] == paid, eid
        assert p["totals"]["gross_earnings"] == gross, eid
        assert _lines(p, "deductions") == {k: D(v) for k, v in ded.items()}, (
            eid
        )  # also proves "no PT in June", "no ESI above 20,000"
        assert p["totals"]["net_pay"] == net, eid
        assert D(p["totals"]["total_deductions"]) == sum(D(v) for v in ded.values()), eid

    # the independent oracle agrees with the figures above
    assert oracle_pf(D(19000)) == 1800 and oracle_pf(D(20000)) == 1800 and oracle_pf(D(50000)) == 1800
    low = await _admin_payslip(client, admin, run["id"], "LOW001")
    assert D(_lines(low, "deductions")["ESI"]) == oracle_esi(D(18850), "0.75")

    # proration of each earning line for the part-month people
    join = await _admin_payslip(client, admin, run["id"], "JOIN01")
    assert join["lop"]["breakdown"]["not_employed"] == 10.0 and _lines(join, "earnings") == {"BASIC": D(10000), "HRA": D(3333)}
    leav = await _admin_payslip(client, admin, run["id"], "LEAV01")
    assert leav["lop"]["breakdown"]["not_employed"] == 15.0 and _lines(leav, "earnings") == {"BASIC": D(20000), "HRA": D(5000)}
    emp = await _admin_payslip(client, admin, run["id"], "EMP001")
    assert emp["lop"]["override"] == 1.5 and "LOP_OVERRIDDEN" in emp["warnings"]
    assert _lines(emp, "earnings") == {"BASIC": D(19000), "HRA": D(7600), "ARREARS": D(2500)}

    # run totals (draft)
    r = await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=admin)
    assert r.json()["totals"] == {
        "gross_earnings": "176283.00",
        "total_deductions": "32269.00",  # 19,582 employee-side + 12,687 employer share recovered (11,640 PF + 1,047 ESI)
        "net_pay": "144014.00",
        "employer_pf": "11640.00",
        "employer_esi": "1047.00",
        "employer_recovered": "12687.00",
        "employer_cost": "0.00",  # remitted 12,687 - recovered 12,687: no double counting
    }

    # process + publish
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=admin)
    assert r.status_code == 200 and r.json()["counts"]["processed"] == 6
    r = await publish_run(client, admin, run["id"])
    assert r.status_code == 200 and r.json()["status"] == "published" and r.json()["counts"]["published"] == 6

    # the register workbook has a row per person and the same totals
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/register", json={}, headers=admin)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Register", "Summary"]
    ws = wb["Register"]
    head = [c.value for c in ws[1]]
    rows = {c[0].value: {h: x.value for h, x in zip(head, c, strict=True)} for c in ws.iter_rows(min_row=2) if c[0].value}
    assert set(rows) == set(expect)
    assert sum(r_["Net Pay"] for r_ in rows.values()) == 144014
    assert rows["LOW001"]["ESI"] == 142 and rows["LOW001"]["Employer ESI"] == 613 and rows["MGR001"]["CANTEEN"] == 500
    assert rows["LOW001"]["ESI Employer Share Recovered"] == 613 and rows["LOW001"]["Employer Cost (net of recovery)"] == 0
    summary = {a.value: b.value for a, b in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Total Employer Cost (net of recovery)"] == 0 and summary["Total PF Employer Share Recovered"] == 11640
    assert summary["Employees"] == 6 and summary["Total Gross"] == 176283


async def test_employee_sees_only_their_own_published_payslips(client, H, world):
    admin = H["admin"]
    run = await new_run(client, admin)
    ids = {e: (await row_for(client, admin, run["id"], e))["id"] for e in ("EMP001", "MGR001")}
    # nothing is visible before publication, not even by id
    assert (await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"] == []
    assert (await client.get(f"/api/payslips/{ids['EMP001']}", headers=H["employee"])).status_code == 404
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=admin)
    assert (await publish_run(client, admin, run["id"])).status_code == 200

    mine = (await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"]
    assert [i["id"] for i in mine] == [ids["EMP001"]]
    assert (await client.get(f"/api/payslips/{ids['EMP001']}", headers=H["employee"])).status_code == 200
    # another employee's payslip: 404 (not 403, so existence is not revealed), PDF too
    for path in (f"/api/payslips/{ids['MGR001']}", f"/api/payslips/{ids['MGR001']}/pdf"):
        r = await client.get(path, headers=H["employee"])
        assert r.status_code == 404 and r.json()["error"]["code"] == "NOT_FOUND", path
    # a manager is an employee for payroll: own only
    assert (await client.get(f"/api/payslips/{ids['EMP001']}", headers=H["manager"])).status_code == 404
    mgr = (await client.get("/api/payslips", params={"year": 2025}, headers=H["manager"])).json()["items"]
    assert [i["id"] for i in mgr] == [ids["MGR001"]]
    # the employee has no way into the admin payroll API
    for path in (f"/api/admin/payroll/payslips/{ids['EMP001']}", f"/api/admin/payroll/runs/{run['id']}"):
        assert (await client.get(path, headers=H["employee"])).status_code == 403, path


async def test_hr_cannot_edit_own_payslip_but_can_edit_colleagues(client, H, world):
    hr = H["hr"]
    run = await new_run(client, hr)
    own = await row_for(client, hr, run["id"], "HR001")
    other = await row_for(client, hr, run["id"], "EMP001")
    url = f"/api/admin/payroll/payslips/{own['id']}"
    attempts = [
        ("PUT", url, {"expected_version": 1, "adjustments": [{"code": "BONUS", "label": "Self bonus", "amount": "99999"}]}),
        ("PUT", url, {"expected_version": 1, "tds": "0", "hold": True}),
        ("DELETE", url, None),
        ("POST", f"{url}/process", None),
    ]  # fmt: skip
    for method, path, body in attempts:
        r = await client.request(method, path, json=body, headers=hr)
        assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN", (method, path)
    # nothing changed
    p = (await client.get(url, headers=hr)).json()
    assert p["version"] == 1 and p["totals"]["net_pay"] == "46400.00" and p["status"] == "draft"
    # own structure and payroll details are also off limits
    r = await client.post(f"/api/admin/payroll/employees/{world['HR001']}/structures",
                          json={"effective_from": "2025-03-01", "basic": "90000", "hra": "0", "allowances": [], "fixed_deductions": [],
                                "tds_monthly": "0", "pf_applicable": True, "pf_on_full_basic": False, "esi_applicable": True,
                                "pt_applicable": True}, headers=hr)  # fmt: skip
    assert r.status_code == 403
    r = await client.put(f"/api/admin/payroll/employees/{world['HR001']}/details",
                         json={"bank_account": "123456789012", "current_password": PASSWORD}, headers=hr)  # fmt: skip
    assert r.status_code == 403
    # a colleague's payslip is fine
    ok = await client.put(
        f"/api/admin/payroll/payslips/{other['id']}", json={"expected_version": 1, "note": "checked"}, headers=hr
    )
    assert ok.status_code == 200
    # only a different person can change HR's payslip
    r = await client.put(url, json={"expected_version": 1, "note": "by admin"}, headers=H["admin"])
    assert r.status_code == 200


async def test_revise_and_republish_after_publication(client, H, world):
    admin, hr = H["admin"], H["hr"]
    run = await new_run(client, admin)
    emp = await _admin_payslip(client, admin, run["id"], "EMP001")
    r = await client.put(
        f"/api/admin/payroll/payslips/{emp['id']}",
        json={"expected_version": emp["version"], "lop_override": 1.5, "lop_override_reason": "Unpaid leave, later approved as paid",
              "adjustments": [{"code": "ARREARS", "label": "May increment arrears", "amount": "2500"}]},
        headers=admin,
    )  # fmt: skip
    assert r.status_code == 200
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=admin)
    assert (await publish_run(client, admin, run["id"])).status_code == 200
    first = (await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"]
    assert len(first) == 1 and first[0]["revision"] == 1 and first[0]["revised"] is False
    old_id = first[0]["id"]
    assert (await client.get(f"/api/payslips/{old_id}", headers=H["employee"])).json()["totals"]["net_pay"] == "24300.00"

    # the published run is locked for normal edits
    r = await client.put(f"/api/admin/payroll/payslips/{old_id}", json={"expected_version": 1, "note": "x"}, headers=admin)
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=admin)
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/reopen", headers=admin)
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"

    # revise: the leave was approved as paid, so there is no LOP any more
    r = await client.post(
        f"/api/admin/payroll/runs/{run['id']}/employees/{world['EMP001']}/revise",
        json={"reason": "Leave approved as paid after payroll"},
        headers=hr,
    )
    assert r.status_code == 201, r.text
    rev = r.json()
    assert rev["revision"] == 2 and rev["status"] == "draft"
    assert [a["code"] for a in rev["inputs"]["adjustments"]] == ["ARREARS"]  # inputs are kept
    r = await client.put(
        f"/api/admin/payroll/payslips/{rev['id']}",
        json={"expected_version": rev["version"], "lop_override": 0, "lop_override_reason": "Leave approved as paid"},
        headers=hr,
    )
    assert r.status_code == 200, r.text
    # 28,000 + 2,500 arrears = 30,500; PF 1,800 + employer PF share 1,800; TDS 1,200  ->  25,700 (hand worked)
    assert r.json()["totals"] == {"gross_earnings": "30500.00", "total_deductions": "4800.00", "net_pay": "25700.00",
                                  "net_pay_words": "Rupees Twenty-Five Thousand Seven Hundred Only"}  # fmt: skip
    # while the revision is a draft the employee still sees revision 1 with the old amounts
    mid = (await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"]
    assert [(i["id"], i["revision"]) for i in mid] == [(old_id, 1)]
    assert (await client.post(f"/api/admin/payroll/payslips/{rev['id']}/process", headers=hr)).status_code == 200
    r = await client.post(f"/api/admin/payroll/payslips/{rev['id']}/publish", json={"current_password": PW}, headers=hr)
    assert r.status_code == 200 and r.json()["status"] == "published"

    live = (await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"]
    assert [(i["id"], i["revision"], i["revised"]) for i in live] == [(rev["id"], 2, True)]
    new = (await client.get(f"/api/payslips/{rev['id']}", headers=H["employee"])).json()
    assert new["totals"]["net_pay"] == "25700.00" and new["revision_info"]["supersedes_id"] == old_id
    old = (await client.get(f"/api/payslips/{old_id}", headers=H["employee"])).json()
    assert old["status"] == "superseded" and old["totals"]["net_pay"] == "24300.00"
    assert old["revision_info"]["superseded_by_id"] == rev["id"]
    # live copy only: 144,014 + 1,000 (no advance recovery in this test) + (25,700 - 24,300)
    r = await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=admin)
    assert r.json()["totals"]["net_pay"] == "146414.00" and r.json()["counts"]["open_revisions"] == 0
    # the revised employee's own list is still one entry per month
    assert len((await client.get("/api/payslips", params={"year": 2025}, headers=H["employee"])).json()["items"]) == 1
    # another employee's payslip was untouched by the revision
    mgr = (await client.get("/api/payslips", params={"year": 2025}, headers=H["manager"])).json()["items"]
    assert len(mgr) == 1 and mgr[0]["revision"] == 1 and PERIOD == "2025-06"


async def test_unauthenticated_and_wrong_role_get_standard_error_envelope(client, H, world):
    r = await client.get("/api/payslips")
    assert r.status_code == 401
    err = r.json()["error"]
    assert set(err) == {"code", "message", "details", "request_id"}
    r = await client.put("/api/admin/payroll/settings", json={"settings": {"payroll.rounding": "paise"}}, headers=H["employee"])
    assert r.status_code == 403 and set(r.json()["error"]) == {"code", "message", "details", "request_id"}
    # the settings body must be wrapped in {"settings": ...}
    r = await client.put("/api/admin/payroll/settings", json={"payroll.rounding": "paise"}, headers=H["hr"])
    assert r.status_code == 422
