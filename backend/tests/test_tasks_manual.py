"""Plan 14: tasks with calendar reminders, and manual attendance by HR/Admin.
2026-10-21 is a Wednesday; office timezone Asia/Kolkata (UTC+5:30); shift 09:30-18:30, 10 min grace."""

from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Attendance, AuditLog, LateRecord, LeaveRequest, LeaveType, Task, User
from app.services import tasks as tasks_svc
from app.services.notifications import OUTBOX
from tests.conftest import USERS, bearer

NOW = dt.datetime(2026, 10, 21, 6, 30, tzinfo=dt.UTC)  # 12:00 IST


@pytest.fixture(autouse=True)
def clock(monkeypatch):
    def _set(when):
        for mod in ("tasks", "attendance", "late", "leave"):
            monkeypatch.setattr(f"app.services.{mod}.utcnow", lambda w=when: w)

    _set(NOW)
    OUTBOX.clear()
    return _set


@pytest.fixture
async def H(tokens_for):
    cache: dict = {}

    async def _h(role):
        if role not in cache:
            cache[role] = bearer(await tokens_for(role))
        return cache[role]

    return _h


async def uid(db, role) -> uuid.UUID:
    return (await db.execute(select(User.id).where(User.employee_id == USERS[role]))).scalar_one()


def task(title="Submit report", due="2026-10-22T10:00:00+05:30", remind=30, **kw):
    return {"title": title, "notes": "Q3 numbers", "due_at": due, "remind_minutes_before": remind, **kw}


def code(r):
    return r.json()["error"]["code"]


# ================================================================ tasks
async def test_personal_task_crud_and_calendar_range(client, H, db):
    h = await H("employee")
    r = await client.post("/api/tasks", headers=h, json=task())
    assert r.status_code == 201, r.text
    b = r.json()
    assert b["due_at"] == "2026-10-22T04:30:00Z" and b["remind_at"] == "2026-10-22T04:00:00Z"
    assert b["remind_minutes_before"] == 30 and b["done"] is False and b["assigned"] is False and b["can_edit"] is True
    row = await db.get(Task, uuid.UUID(b["id"]))
    assert row.notes == "Q3 numbers"  # encrypted column, transparent
    lst = (await client.get("/api/tasks", headers=h, params={"from": "2026-10-01", "to": "2026-10-31"})).json()
    assert [t["id"] for t in lst["items"]] == [b["id"]]
    assert (await client.get("/api/tasks", headers=h, params={"from": "2026-11-01", "to": "2026-11-30"})).json()["items"] == []
    assert [t["id"] for t in (await client.get("/api/tasks", headers=h)).json()["items"]] == [b["id"]]  # default: this month
    upd = await client.put(f"/api/tasks/{b['id']}", headers=h, json=task(title="Submit final report", remind=None))
    assert upd.status_code == 200 and upd.json()["title"] == "Submit final report" and upd.json()["remind_at"] is None
    d = await client.post(f"/api/tasks/{b['id']}/done", headers=h, json={"done": True})
    assert d.json()["done"] is True
    assert (await client.delete(f"/api/tasks/{b['id']}", headers=h)).status_code == 204
    assert (await client.get(f"/api/tasks/{b['id']}", headers=h)).status_code == 404


@pytest.mark.parametrize(
    "bad",
    [{"title": ""}, {"due_at": "2026-10-22T10:00:00"}, {"remind_minutes_before": -1}, {"remind_minutes_before": 20000}, {"x": 1}],
)
async def test_task_validation(client, H, bad):
    r = await client.post("/api/tasks", headers=await H("employee"), json={**task(), **bad})
    assert r.status_code == 422


async def test_range_limit(client, H):
    r = await client.get("/api/tasks", headers=await H("employee"), params={"from": "2026-01-01", "to": "2026-12-31"})
    assert r.status_code == 422 and code(r) == "INVALID_DATE_RANGE"


async def test_other_people_cannot_see_or_change_a_task(client, H):
    r = await client.post("/api/tasks", headers=await H("employee"), json=task())
    other = await H("hr")
    tid = r.json()["id"]
    assert (await client.get(f"/api/tasks/{tid}", headers=other)).status_code == 404
    assert (await client.put(f"/api/tasks/{tid}", headers=other, json=task())).status_code == 404
    assert (await client.post(f"/api/tasks/{tid}/done", headers=other, json={"done": True})).status_code == 404
    assert (await client.delete(f"/api/tasks/{tid}", headers=other)).status_code == 404


async def test_manager_assigns_to_direct_report_only(client, H, db):
    mgr = await H("manager")
    emp_id = await uid(db, "employee")
    r = await client.post("/api/tasks", headers=mgr, json=task(user_id=str(emp_id)))
    assert r.status_code == 201, r.text
    assert r.json()["assigned"] is True and r.json()["user"]["employee_id"] == USERS["employee"]
    assert OUTBOX[-1]["payload"]["type"] == "task_assigned" and OUTBOX[-1]["user_ids"] == [str(emp_id)]
    hr_id = await uid(db, "hr")  # not a direct report of the manager
    r2 = await client.post("/api/tasks", headers=mgr, json=task(user_id=str(hr_id)))
    assert r2.status_code == 422 and code(r2) == "VALIDATION_ERROR"
    # the creator sees it under assigned_by_me, not in their own list
    assert (await client.get("/api/tasks", headers=mgr)).json()["items"] == []
    mine = (await client.get("/api/tasks", headers=mgr, params={"assigned_by_me": True})).json()["items"]
    assert [t["id"] for t in mine] == [r.json()["id"]]
    # the assignee may complete it but not edit or delete it
    emp = await H("employee")
    tid = r.json()["id"]
    got = (await client.get(f"/api/tasks/{tid}", headers=emp)).json()
    assert got["can_edit"] is False
    assert (await client.put(f"/api/tasks/{tid}", headers=emp, json=task())).status_code == 403
    assert (await client.delete(f"/api/tasks/{tid}", headers=emp)).status_code == 403
    OUTBOX.clear()
    assert (await client.post(f"/api/tasks/{tid}/done", headers=emp, json={"done": True})).json()["done"] is True
    assert OUTBOX[-1]["payload"]["type"] == "task_done" and OUTBOX[-1]["user_ids"] == [str(await uid(db, "manager"))]


async def test_employee_cannot_assign_and_hr_can_assign_anyone(client, H, db):
    r = await client.post("/api/tasks", headers=await H("employee"), json=task(user_id=str(await uid(db, "hr"))))
    assert r.status_code == 403
    r = await client.post("/api/tasks", headers=await H("hr"), json=task(user_id=str(await uid(db, "manager"))))
    assert r.status_code == 201


async def test_reminder_job_pushes_once_and_skips_done_or_past(client, H, db, clock):
    h = await H("employee")
    a = (await client.post("/api/tasks", headers=h, json=task(title="A", remind=30))).json()  # remind 09:30 IST tomorrow
    await client.post("/api/tasks", headers=h, json=task(title="B", remind=0))
    c = (await client.post("/api/tasks", headers=h, json=task(title="C", due="2026-10-21T11:00:00+05:30", remind=0))).json()
    assert c["remind_at"] is not None  # reminder time already past at save: armed but marked sent
    done = (await client.post("/api/tasks", headers=h, json=task(title="D", remind=0))).json()
    await client.post(f"/api/tasks/{done['id']}/done", headers=h, json={"done": True})
    OUTBOX.clear()
    assert await tasks_svc.send_due_reminders(db, NOW) == 0
    later = dt.datetime(2026, 10, 22, 4, 0, tzinfo=dt.UTC)  # 09:30 IST
    assert await tasks_svc.send_due_reminders(db, later) == 1
    assert OUTBOX[-1]["payload"]["type"] == "task_reminder" and OUTBOX[-1]["payload"]["title"] == "Reminder: A"
    assert OUTBOX[-1]["payload"]["body"] == "Due 22 Oct, 10:00"
    assert await tasks_svc.send_due_reminders(db, later) == 0  # once only
    assert await tasks_svc.send_due_reminders(db, later + dt.timedelta(minutes=30)) == 1  # B at the due time
    assert {x["payload"]["title"] for x in OUTBOX} == {"Reminder: A", "Reminder: B"}
    # moving the due time re-arms the reminder
    clock(later + dt.timedelta(hours=1))
    r = await client.put(f"/api/tasks/{a['id']}", headers=h, json=task(title="A", due="2026-10-23T10:00:00+05:30", remind=30))
    assert r.status_code == 200
    assert await tasks_svc.send_due_reminders(db, dt.datetime(2026, 10, 23, 4, 0, tzinfo=dt.UTC)) == 1


# ================================================================ manual attendance
def manual(cin="09:20", cout="18:40", reason="Forgot to punch, confirmed by manager"):
    return {"check_in": cin, "check_out": cout, "reason": reason}


async def test_hr_enters_manual_attendance_and_late_is_derived(client, H, db):
    hr = await H("hr")
    emp = await uid(db, "employee")
    url = f"/api/admin/attendance/{emp}/2026-10-20"
    assert (await client.get(url, headers=hr)).json()["record"] is None
    r = await client.put(url, headers=hr, json=manual())
    assert r.status_code == 200, r.text
    rec = r.json()["record"]
    assert rec["manual"] is True and rec["status"] == "present" and rec["minutes_late"] == 0
    assert rec["worked_minutes"] == 560 and rec["manual_reason"] == "Forgot to punch, confirmed by manager"
    assert OUTBOX[-1]["payload"]["type"] == "attendance_manual" and OUTBOX[-1]["user_ids"] == [str(emp)]
    # correction: now late -> late record rebuilt; the employee sees it as manual in history
    r = await client.put(url, headers=hr, json=manual(cin="10:05", cout="18:40"))
    assert r.json()["record"]["status"] == "late" and r.json()["record"]["minutes_late"] == 35
    lates = (await db.execute(select(LateRecord).where(LateRecord.user_id == emp))).scalars().all()
    assert [lr.minutes_late for lr in lates] == [35]
    hist = (await client.get("/api/attendance", headers=await H("employee"))).json()["items"]
    assert hist[0]["manual"] is True and "manual_reason" not in hist[0]
    logs = (await db.execute(select(AuditLog).where(AuditLog.action == "attendance.manual"))).scalars().all()
    assert len(logs) == 2 and logs[1].before["status"] == "present" and logs[1].after["status"] == "late"
    assert "Forgot" not in repr(logs[1].after)  # the reason text is not copied into the audit log


async def test_short_day_is_half_day_and_open_day_allowed(client, H, db):
    hr = await H("hr")
    emp = await uid(db, "employee")
    r = await client.put(f"/api/admin/attendance/{emp}/2026-10-19", headers=hr, json=manual(cin="09:30", cout="12:00"))
    assert r.json()["record"]["status"] == "half_day"
    r = await client.put(f"/api/admin/attendance/{emp}/2026-10-21", headers=hr, json=manual(cin="09:25", cout=None))
    assert r.status_code == 200 and r.json()["record"]["check_out_at"] is None


async def test_manual_corrects_a_real_punch_but_cannot_delete_it(client, H, db, clock):
    emp_h = await H("employee")
    clock(dt.datetime(2026, 10, 21, 4, 45, tzinfo=dt.UTC))  # 10:15 IST: late
    fix = {"lat": 12.9716, "lng": 77.5946, "accuracy_m": 10, "fix_time": "2026-10-21T04:45:00Z"}
    assert (await client.post("/api/attendance/check-in", headers=emp_h, json=fix)).status_code == 201
    clock(NOW)
    emp = await uid(db, "employee")
    url = f"/api/admin/attendance/{emp}/2026-10-21"
    r = await client.put(url, headers=await H("hr"), json=manual(cin="09:30", cout="11:59", reason="Biometric was down at gate"))
    assert r.status_code == 200 and r.json()["record"]["minutes_late"] == 0
    row = (await db.execute(select(Attendance).where(Attendance.user_id == emp))).scalar_one()
    assert row.check_in_within_fence is True  # punch evidence kept
    r = await client.delete(url, headers=await H("hr"))
    assert r.status_code == 409


async def test_delete_manual_entry(client, H, db):
    hr = await H("hr")
    emp = await uid(db, "employee")
    url = f"/api/admin/attendance/{emp}/2026-10-20"
    await client.put(url, headers=hr, json=manual(cin="10:30"))
    assert (await client.delete(url, headers=hr)).status_code == 204
    assert (await client.get(url, headers=hr)).json()["record"] is None
    assert (await db.execute(select(LateRecord))).scalars().all() == []
    assert (await client.delete(url, headers=hr)).status_code == 404


@pytest.mark.parametrize(
    ("day", "body", "status", "err"),
    [
        ("2026-10-22", manual(), 422, "INVALID_DATE_RANGE"),  # future date
        ("2026-10-20", manual(cin="18:00", cout="09:00"), 422, "VALIDATION_ERROR"),  # out before in (day shift)
        ("2026-10-21", manual(cin="09:30", cout="13:00"), 422, "VALIDATION_ERROR"),  # check-out later than now (12:00)
        ("2026-10-20", manual(reason=" "), 422, "VALIDATION_ERROR"),
    ],
)
async def test_manual_validation(client, H, db, day, body, status, err):
    emp = await uid(db, "employee")
    r = await client.put(f"/api/admin/attendance/{emp}/{day}", headers=await H("hr"), json=body)
    assert r.status_code == status and code(r) == err, r.text


async def test_manual_permissions_and_no_self(client, H, db):
    emp = await uid(db, "employee")
    for role in ("employee", "manager"):
        r = await client.put(f"/api/admin/attendance/{emp}/2026-10-20", headers=await H(role), json=manual())
        assert r.status_code == 403
    hr_id = await uid(db, "hr")
    r = await client.put(f"/api/admin/attendance/{hr_id}/2026-10-20", headers=await H("hr"), json=manual())
    assert r.status_code == 403
    r = await client.put(f"/api/admin/attendance/{uuid.uuid4()}/2026-10-20", headers=await H("hr"), json=manual())
    assert r.status_code == 404


async def test_manual_refused_on_full_day_leave(client, H, db):
    emp = await uid(db, "employee")
    lt = LeaveType(code="CL", name="Casual Leave", annual_quota=12)
    db.add(lt)
    await db.flush()
    db.add(LeaveRequest(user_id=emp, leave_type_id=lt.id, from_date=dt.date(2026, 10, 20), to_date=dt.date(2026, 10, 20),
                        days=1, reason="x", status="approved"))  # fmt: skip
    await db.commit()
    r = await client.put(f"/api/admin/attendance/{emp}/2026-10-20", headers=await H("hr"), json=manual())
    assert r.status_code == 409 and code(r) == "ON_LEAVE"


async def test_manual_entry_in_export(client, H, db):
    emp = await uid(db, "employee")
    await client.put(f"/api/admin/attendance/{emp}/2026-10-20", headers=await H("hr"), json=manual())
    r = await client.get(
        "/api/admin/export", headers=await H("admin"), params={"type": "attendance", "from": "2026-10-20", "to": "2026-10-20"}
    )
    ws = load_workbook(io.BytesIO(r.content))["Attendance"]
    head = [c.value for c in ws[1]]
    row = [c.value for c in next(ws.iter_rows(min_row=2))]
    assert row[head.index("Manual Entry")] == "Yes"
