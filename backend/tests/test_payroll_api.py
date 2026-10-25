"""Payroll API (contract section 13): settings, structures, details, runs, lifecycle, revisions, RBAC, self rule, audit."""

import datetime as dt

import pytest
from sqlalchemy import select, text

from app.models import AuditLog, EmployeePayrollDetails, PayrollRun, Payslip
from app.services import notifications
from tests.conftest import PASSWORD, bearer
from tests.payroll_fixtures import (
    PERIOD,
    add_structure,
    exclude_from_payroll,
    new_run,
    payslips_of,
    process_and_publish,
    publish_run,
    row_for,
    set_setting,
    structure_body,
)

# ================================================================================================ RBAC
ADMIN_ROUTES = [
    ("GET", "/api/admin/payroll/settings"),
    ("PUT", "/api/admin/payroll/settings"),
    ("GET", "/api/admin/payroll/employees"),
    ("GET", "/api/admin/payroll/employees/{u}/structures"),
    ("POST", "/api/admin/payroll/employees/{u}/structures"),
    ("GET", "/api/admin/payroll/employees/{u}/details"),
    ("PUT", "/api/admin/payroll/employees/{u}/details"),
    ("GET", "/api/admin/payroll/runs"),
    ("POST", "/api/admin/payroll/runs"),
    ("GET", "/api/admin/payroll/runs/{u}"),
    ("POST", "/api/admin/payroll/runs/{u}/refresh"),
    ("POST", "/api/admin/payroll/runs/{u}/process"),
    ("POST", "/api/admin/payroll/runs/{u}/publish"),
    ("GET", "/api/admin/payroll/runs/{u}/payslips"),
    ("POST", "/api/admin/payroll/runs/{u}/register"),
    ("GET", "/api/admin/payroll/payslips/{u}"),
    ("PUT", "/api/admin/payroll/payslips/{u}"),
    ("GET", "/api/admin/payroll/payslips/{u}/pdf"),
    ("DELETE", "/api/admin/payroll/structures/{u}"),
]


@pytest.mark.parametrize("role", ["employee", "manager"])
@pytest.mark.parametrize("method,path", ADMIN_ROUTES)
async def test_non_hr_roles_cannot_use_admin_payroll_routes(client, H, uid, role, method, path):
    r = await client.request(
        method, path.format(u=uid["employee"]), headers=H[role], json={} if method in ("POST", "PUT") else None
    )
    assert r.status_code == 403, (path, r.text)
    assert r.json()["error"]["code"] == "FORBIDDEN"


@pytest.mark.parametrize("method,path", ADMIN_ROUTES[:3] + [("GET", "/api/payslips"), ("GET", "/api/payslips/{u}")])
async def test_unauthenticated_is_401(client, uid, method, path):
    r = await client.request(method, path.format(u=uid["employee"]))
    assert r.status_code == 401


async def test_permission_catalogue_matches_contract(client, H):
    me = await client.get("/api/auth/me", headers=H["employee"])
    assert "payroll.view_own" in me.json()["permissions"]
    hr = (await client.get("/api/auth/me", headers=H["hr"])).json()["permissions"]
    assert {"payroll.view_own", "payroll.view_all", "payroll.manage", "payroll.publish"} <= set(hr)


# ================================================================================================ settings
async def test_settings_defaults_and_merge_update(client, H, db):
    r = await client.get("/api/admin/payroll/settings", headers=H["hr"])
    s = r.json()["settings"]
    assert s["payroll.pt_frequency"] == "half_yearly" and s["payroll.pf_wage_ceiling"] == 15000
    assert s["payroll.pt_slabs"][0]["months"] == [3, 9]
    r = await client.put(
        "/api/admin/payroll/settings",
        json={"settings": {"payroll.pf_wage_ceiling": 20000, "payroll.rounding": "paise"}},
        headers=H["hr"],
    )
    assert r.status_code == 200
    s = r.json()["settings"]
    assert (
        s["payroll.pf_wage_ceiling"] == 20000 and s["payroll.rounding"] == "paise" and s["payroll.pt_frequency"] == "half_yearly"
    )
    row = (await db.execute(select(AuditLog).where(AuditLog.action == "payroll.settings.update"))).scalar_one()
    assert {c["key"] for c in row.after["changes"]} == {"payroll.pf_wage_ceiling", "payroll.rounding"}


@pytest.mark.parametrize(
    "key,value",
    [
        ("payroll.pt_frequency", "weekly"), ("payroll.pay_days_basis", "x"), ("payroll.pf_employee_rate_pct", 21),
        ("payroll.pf_employee_rate_pct", 12.345), ("payroll.pf_wage_ceiling", -1), ("payroll.esi_employee_rate_pct", 11),
        ("payroll.mask_pan", "yes"), ("payroll.rounding", "euro"), ("payroll.company_address", "x" * 301),
        ("payroll.payslip_footer", "x" * 201), ("payroll.nope", 1), ("payroll.pf_wage_ceiling", True),
        ("payroll.pt_slabs", "none"),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": "10", "amount": "1", "months": [13]}]),
        ("payroll.pt_slabs", [{"min_gross": "10", "max_gross": "5", "amount": "1", "months": None}]),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": "100", "amount": "1", "months": None},
                              {"min_gross": "50", "max_gross": None, "amount": "2", "months": None}]),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": "100", "amount": "1", "months": [3]},
                              {"min_gross": "100", "max_gross": None, "amount": "2", "months": [3, 4]}]),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": None, "amount": 1.5, "months": None}]),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": None, "amount": "1", "months": None, "x": 1}]),
        ("payroll.pt_slabs", [{"min_gross": "0", "max_gross": None, "amount": "1", "months": None}] * 21),
    ],
)  # fmt: skip
async def test_settings_validation_rejects_bad_values(client, H, key, value):
    r = await client.put("/api/admin/payroll/settings", json={"settings": {key: value}}, headers=H["hr"])
    assert r.status_code == 422, r.text
    assert r.json()["error"]["code"] == "VALIDATION_ERROR"
    assert r.json()["error"]["details"][0]["field"] == key


async def test_pt_slab_with_month_specific_override_is_allowed(client, H):
    slabs = [{"min_gross": "25000.00", "max_gross": None, "amount": "200.00", "months": None},
             {"min_gross": "25000.00", "max_gross": None, "amount": "300.00", "months": [2]}]  # fmt: skip
    r = await client.put("/api/admin/payroll/settings", json={"settings": {"payroll.pt_slabs": slabs}}, headers=H["hr"])
    assert r.status_code == 200 and r.json()["settings"]["payroll.pt_slabs"] == slabs
    r = await client.put("/api/admin/payroll/settings", json={"settings": {"payroll.pt_slabs": []}}, headers=H["hr"])
    assert r.status_code == 200 and r.json()["settings"]["payroll.pt_slabs"] == []


async def test_payroll_keys_are_not_accepted_by_the_generic_settings_endpoint(client, H):
    r = await client.put("/api/admin/settings", json={"settings": {"payroll.pt_frequency": "monthly"}}, headers=H["admin"])
    assert r.status_code == 422
    g = (await client.get("/api/admin/settings", headers=H["admin"])).json()["settings"]
    assert not [k for k in g if k.startswith("payroll.")]


async def test_settings_change_does_not_alter_an_existing_run_until_refresh(client, payroll):
    run = await new_run(client, payroll["hr"])
    assert run["config"]["pf_wage_ceiling"] == 15000
    await set_setting(client, payroll["hr"], pf_wage_ceiling=0)
    r = await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])
    assert r.json()["config"]["pf_wage_ceiling"] == 15000
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/refresh", headers=payroll["hr"])
    assert r.json()["config"]["pf_wage_ceiling"] == 0


# ================================================================================================ structures
async def test_structure_crud_effective_to_and_in_use(client, H, uid):
    a = await add_structure(client, H["hr"], uid["employee"], basic="20000", hra="5000", eff="2025-01-01",
                            allowances=[{"code": "SPECIAL", "name": "Special", "amount": "3000.50", "prorate": True, "esi_wage": True}],
                            fixed_deductions=[{"code": "CANTEEN", "name": "Canteen", "amount": "500"}], tds_monthly="1200")  # fmt: skip
    assert a["monthly_gross"] == "28000.50" and a["effective_to"] is None and a["in_use"] is False and a["basic"] == "20000.00"
    b = await add_structure(client, H["hr"], uid["employee"], basic="22000", eff="2025-07-01")
    r = await client.get(f"/api/admin/payroll/employees/{uid['employee']}/structures", headers=H["hr"])
    items = r.json()["items"]
    assert [i["effective_from"] for i in items] == ["2025-07-01", "2025-01-01"]
    assert items[1]["effective_to"] == "2025-06-30" and items[0]["effective_to"] is None
    r = await client.put(
        f"/api/admin/payroll/structures/{b['id']}", json=structure_body(basic="23000", eff="2025-07-01"), headers=H["hr"]
    )
    assert r.status_code == 200 and r.json()["basic"] == "23000.00"
    r = await client.delete(f"/api/admin/payroll/structures/{b['id']}", headers=H["hr"])
    assert r.status_code == 204


async def test_structure_validation(client, H, uid):
    url = f"/api/admin/payroll/employees/{uid['employee']}/structures"
    r = await client.post(url, json=structure_body(eff="2025-01-15"), headers=H["hr"])
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_PERIOD"
    await add_structure(client, H["hr"], uid["employee"])
    r = await client.post(url, json=structure_body(), headers=H["hr"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "STRUCTURE_EXISTS"
    bad = [
        structure_body(eff="2025-02-01", basic=20000),  # money must be a string
        structure_body(eff="2025-02-01", basic="20000.123"),
        structure_body(eff="2025-02-01", basic="0"),
        structure_body(eff="2025-02-01", basic="10000001"),
        structure_body(eff="2025-02-01", allowances=[{"code": "PF", "name": "x", "amount": "1"}]),
        structure_body(eff="2025-02-01", allowances=[{"code": "a b", "name": "x", "amount": "1"}]),
        structure_body(
            eff="2025-02-01", allowances=[{"code": "A1", "name": "x", "amount": "1"}, {"code": "A1", "name": "y", "amount": "1"}]
        ),
        structure_body(eff="2025-02-01", fixed_deductions=[{"code": "BONUS", "name": "x", "amount": "1"}]),
        structure_body(eff="2025-02-01", allowances=[{"code": f"A{i}", "name": "x", "amount": "1"} for i in range(16)]),
        structure_body(eff="2025-02-01", unknown_field=1),
    ]
    for body in bad:
        r = await client.post(url, json=body, headers=H["hr"])
        assert r.status_code == 422, body
        assert "20000.123" not in r.text  # the submitted value is never echoed


async def test_cannot_edit_own_structure(client, H, uid):
    r = await client.post(f"/api/admin/payroll/employees/{uid['hr']}/structures", json=structure_body(), headers=H["hr"])
    assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN"
    s = await add_structure(client, H["admin"], uid["hr"])  # another HR/admin may do it
    r = await client.put(f"/api/admin/payroll/structures/{s['id']}", json=structure_body(basic="99999"), headers=H["hr"])
    assert r.status_code == 403
    r = await client.delete(f"/api/admin/payroll/structures/{s['id']}", headers=H["hr"])
    assert r.status_code == 403


async def test_hr_may_set_admin_and_manager_salary_privilege_rules_do_not_apply(client, H, uid):
    await add_structure(client, H["hr"], uid["admin"])
    await add_structure(client, H["hr"], uid["manager"])


async def test_unknown_employee_is_404(client, H):
    r = await client.get("/api/admin/payroll/employees/00000000-0000-0000-0000-000000000000/structures", headers=H["hr"])
    assert r.status_code == 404


async def test_employee_list_flags_and_filters(client, H, uid):
    await add_structure(client, H["hr"], uid["employee"], basic="20000", hra="5000")
    r = await client.get("/api/admin/payroll/employees", params={"has_structure": "true"}, headers=H["hr"])
    items = r.json()["items"]
    assert [i["user"]["employee_id"] for i in items] == ["EMP001"]
    assert items[0]["current_structure"]["monthly_gross"] == "25000.00" and set(items[0]["flags"]) == {
        "NO_BANK_DETAILS",
        "NO_PAN",
    }
    r = await client.get("/api/admin/payroll/employees", params={"has_structure": "false", "q": "manager"}, headers=H["hr"])
    assert [i["user"]["employee_id"] for i in r.json()["items"]] == ["MGR001"]
    assert r.json()["total"] == 1 and r.json()["page"] == 1


# ================================================================================================ details
async def test_details_masked_reveal_and_audit(client, H, uid, db):
    body = {"designation": "Analyst", "payment_mode": "bank_transfer", "bank_name": "HDFC Bank", "bank_account": "1234 5678 9012",
            "bank_ifsc": "hdfc0001234", "pan": "abcde1234f", "uan": "100200300400", "esi_ip_number": "1234567890", "current_password": PASSWORD}  # fmt: skip
    r = await client.put(f"/api/admin/payroll/employees/{uid['employee']}/details", json=body, headers=H["hr"])
    assert r.status_code == 200, r.text
    d = r.json()
    assert (
        d["bank_account"] == "XXXXXXXX9012"
        and d["pan"] == "XXXXX1234X"
        and d["uan"] == "XXXXXXXX0400"
        and d["esi_ip_number"] == "XXXXXX7890"
    )
    assert d["bank_ifsc"] == "HDFC0001234" and d["revealed"] is False and d["updated_by"]["employee_id"] == "HR001"
    r = await client.get(f"/api/admin/payroll/employees/{uid['employee']}/details", params={"reveal": "true"}, headers=H["hr"])
    d = r.json()
    assert d["pan"] == "ABCDE1234F" and d["bank_account"] == "123456789012" and d["revealed"] is True
    raw = (await db.execute(text("SELECT bank_account_enc, pan_enc FROM employee_payroll_details"))).one()
    assert b"123456789012" not in bytes(raw[0]) and b"ABCDE1234F" not in bytes(raw[1])  # encrypted at rest
    logs = (await db.execute(select(AuditLog).where(AuditLog.action.like("payroll.details.%")))).scalars().all()
    blob = " ".join(str(x.after) for x in logs)
    assert "ABCDE1234F" not in blob and "123456789012" not in blob and "HDFC" not in blob
    upd = next(x for x in logs if x.action == "payroll.details.update")
    assert set(upd.after["fields"]) == {
        "designation",
        "bank_name",
        "bank_account",
        "bank_ifsc",
        "pan",
        "uan",
        "esi_ip_number",
    }  # payment_mode did not change
    assert any(x.action == "payroll.details.view" and x.after == {"revealed": True} for x in logs)


async def test_details_update_needs_password_and_valid_formats(client, H, uid):
    url = f"/api/admin/payroll/employees/{uid['employee']}/details"
    r = await client.put(url, json={"designation": "X"}, headers=H["hr"])
    assert r.status_code == 422 and r.json()["error"]["details"][0]["field"] == "current_password"
    r = await client.put(url, json={"designation": "X", "current_password": "wrong-password-1"}, headers=H["hr"])
    assert r.status_code == 403 and r.json()["error"]["code"] == "REAUTH_FAILED"
    for bad in ({"pan": "ABC"}, {"uan": "123"}, {"bank_ifsc": "HDFC1001234"}, {"bank_account": "123"}, {"esi_ip_number": "12"},
                {"payment_mode": "bitcoin"}, {"include_in_payroll": None}, {"unknown": 1}):  # fmt: skip
        r = await client.put(url, json={**bad, "current_password": PASSWORD}, headers=H["hr"])
        assert r.status_code == 422, bad


async def test_details_null_clears_and_only_sent_keys_change(client, H, uid):
    url = f"/api/admin/payroll/employees/{uid['employee']}/details"
    await client.put(url, json={"pan": "ABCDE1234F", "designation": "A", "current_password": PASSWORD}, headers=H["hr"])
    r = await client.put(url, json={"pan": None, "current_password": PASSWORD}, headers=H["hr"])
    assert r.json()["pan"] is None and r.json()["designation"] == "A"


async def test_cannot_edit_own_details(client, H, uid):
    r = await client.put(
        f"/api/admin/payroll/employees/{uid['hr']}/details",
        json={"designation": "x", "current_password": PASSWORD},
        headers=H["hr"],
    )
    assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN"


# ================================================================================================ run creation / process
async def test_create_run_drafts_every_eligible_employee(client, payroll, db):
    run = await new_run(client, payroll["hr"])
    assert run["status"] == "draft" and run["period"] == PERIOD and run["counts"]["employees"] == 2
    assert run["attendance_cutoff"] == "2025-06-30" and run["pay_date"] == "2025-06-30"
    items = await payslips_of(client, payroll["hr"], run["id"])
    assert {i["user"]["employee_id"] for i in items} == {"EMP001", "MGR001"}
    emp = next(i for i in items if i["user"]["employee_id"] == "EMP001")
    # 20000 + 5000 for 30 of 30 days; PF 1800 (basic 20000 capped at 15000) + 1800 employer share recovered = 3600; ESI: 25000 > 20000 -> none; PT: June not a PT month
    assert emp["gross_earnings"] == "25000.00" and emp["total_deductions"] == "3600.00" and emp["net_pay"] == "21400.00"
    assert emp["paid_days"] == 30.0 and emp["lop_days"] == 0.0 and emp["errors"] == []
    assert set(emp["warnings"]) == {"NO_BANK_DETAILS", "NO_PAN"}


async def test_create_run_rules(client, payroll):
    r = await client.post("/api/admin/payroll/runs", json={"period": "2999-01"}, headers=payroll["hr"])
    assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_PERIOD"
    for bad in ("2025-13", "2025-6", "abc", "1999-12"):
        r = await client.post("/api/admin/payroll/runs", json={"period": bad}, headers=payroll["hr"])
        assert r.status_code == 422 and r.json()["error"]["code"] == "INVALID_PERIOD", bad
    await new_run(client, payroll["hr"])
    r = await client.post("/api/admin/payroll/runs", json={"period": PERIOD}, headers=payroll["hr"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_RUN_EXISTS"


async def test_missing_structure_blocks_process_with_a_list_and_changes_nothing(client, payroll, uid, db):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    run = await new_run(client, payroll["hr"], period="2024-12")  # before the 2025-01-01 structures
    items = await payslips_of(client, payroll["hr"], run["id"])
    assert items[0]["errors"] == ["STRUCTURE_MISSING"] and items[0]["net_pay"] is None
    assert run["counts"]["errors"] == 1
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_BLOCKED"
    assert r.json()["error"]["details"] == [{"field": "EMP001", "message": "STRUCTURE_MISSING"}]
    got = (await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])).json()
    assert got["status"] == "draft"


async def test_negative_net_blocks_process_until_the_fixed_deduction_is_waived(client, payroll, uid):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    s = await add_structure(client, payroll["hr"], uid["employee"], basic="1000", eff="2025-03-01",
                            fixed_deductions=[{"code": "RENT", "name": "Rent", "amount": "5000"}], pf_applicable=False)  # fmt: skip
    run = await new_run(client, payroll["hr"])
    row = (await payslips_of(client, payroll["hr"], run["id"]))[0]
    assert row["errors"] == ["NEGATIVE_NET_PAY"] and row["net_pay"] is None
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    assert r.status_code == 409 and r.json()["error"]["details"] == [{"field": "EMP001", "message": "NEGATIVE_NET_PAY"}]
    # hold removes it from the process; the rest of the run can go ahead
    got = (await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).json()
    assert got["earnings"] == [] and got["totals"] is None and got["errors"] == ["NEGATIVE_NET_PAY"]
    r = await client.put(
        f"/api/admin/payroll/payslips/{row['id']}", json={"expected_version": got["version"], "hold": True}, headers=payroll["hr"]
    )
    assert r.status_code == 200 and r.json()["hold"] is True
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    assert r.status_code == 200 and r.json()["status"] == "processed" and r.json()["counts"]["held"] == 1
    assert s["id"]


async def test_draft_edit_recomputes_with_version_check_and_adjustments(client, payroll):
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    pid = row["id"]
    got = (await client.get(f"/api/admin/payroll/payslips/{pid}", headers=payroll["hr"])).json()
    assert got["version"] >= 1 and got["inputs"]["tds"] == "0.00" and got["lop"]["auto"] == 0.0
    body = {"expected_version": got["version"], "tds": "500", "pt_override": "200",
            "adjustments": [{"code": "BONUS", "label": "Diwali", "amount": "5000"}, {"code": "LOAN_EMI", "label": "Loan", "amount": "1000"}],
            "note": "checked"}  # fmt: skip
    r = await client.put(f"/api/admin/payroll/payslips/{pid}", json=body, headers=payroll["hr"])
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["totals"]["gross_earnings"] == "30000.00"
    codes = {d["code"]: d["amount"] for d in p["deductions"]}
    assert (
        codes["TDS"] == "500.00"
        and codes["PT"] == "200.00"
        and codes["LOAN_EMI"] == "1000.00"
        and codes["PF"] == "1800.00"
        and codes["PF_ER_REC"] == "1800.00"
    )
    assert p["totals"]["net_pay"] == "24700.00"  # 30000 - (1800+1800+200+500+1000)
    assert p["inputs"]["note"] == "checked" and p["version"] == got["version"] + 1
    # stale version
    r = await client.put(
        f"/api/admin/payroll/payslips/{pid}", json={**body, "expected_version": got["version"]}, headers=payroll["hr"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "VERSION_CONFLICT"
    # only sent keys change; null clears
    r = await client.put(
        f"/api/admin/payroll/payslips/{pid}",
        json={"expected_version": p["version"], "pt_override": None, "tds": None},
        headers=payroll["hr"],
    )
    assert r.status_code == 200
    q = r.json()
    assert q["inputs"]["pt_override"] is None and q["inputs"]["tds"] == "0.00" and len(q["inputs"]["adjustments"]) == 2


async def test_payslip_edit_validation(client, payroll):
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    v = (await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).json()["version"]
    url = f"/api/admin/payroll/payslips/{row['id']}"
    bad = [
        {"tds": 500}, {"tds": "-1"}, {"pt_override": "1.234"}, {"lop_override": 1.0},  # reason missing
        {"lop_override": 1.3, "lop_override_reason": "because"}, {"lop_override": -1, "lop_override_reason": "because"},
        {"lop_override": 1.0, "lop_override_reason": "no"}, {"hold": None}, {"adjustments": None},
        {"adjustments": [{"code": "PF", "label": "x", "amount": "1"}]},
        {"adjustments": [{"code": "BONUS", "label": "", "amount": "1"}]},
        {"adjustments": [{"code": "BONUS", "label": "x", "amount": "1"}] * 11},
        {"note": "x" * 201}, {"surprise": 1},
    ]  # fmt: skip
    for b in bad:
        r = await client.put(url, json={"expected_version": v, **b}, headers=payroll["hr"])
        assert r.status_code == 422, b
    r = await client.put(
        url, json={"expected_version": v, "lop_override": 31, "lop_override_reason": "too many days"}, headers=payroll["hr"]
    )
    assert r.status_code == 422 and r.json()["error"]["details"][0]["field"] == "lop_override"  # above the employed days


async def test_lop_override_in_half_day_steps_with_reason(client, payroll):
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    url = f"/api/admin/payroll/payslips/{row['id']}"
    v = (await client.get(url, headers=payroll["hr"])).json()["version"]
    r = await client.put(
        url,
        json={"expected_version": v, "lop_override": 1.5, "lop_override_reason": "Confirmed absent twice"},
        headers=payroll["hr"],
    )
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["days"] == {"basis_days": 30.0, "employed_days": 30.0, "lop_days": 1.5, "paid_days": 28.5}
    assert (
        p["lop"]["override"] == 1.5
        and p["lop"]["override_reason"] == "Confirmed absent twice"
        and "LOP_OVERRIDDEN" in p["warnings"]
    )
    assert p["earnings"][0]["amount"] == "19000.00"  # 20000 x 28.5 / 30
    r = await client.put(url, json={"expected_version": p["version"], "lop_override": None}, headers=payroll["hr"])
    assert (
        r.json()["days"]["lop_days"] == 0.0
        and "LOP_OVERRIDDEN" not in r.json()["warnings"]
        and r.json()["lop"]["override_reason"] is None
    )


async def test_refresh_keeps_manual_inputs_adds_new_employees_and_flags_ineligible(client, payroll, uid):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    run = await new_run(client, payroll["hr"])
    assert run["counts"]["employees"] == 1
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    url = f"/api/admin/payroll/payslips/{row['id']}"
    v = (await client.get(url, headers=payroll["hr"])).json()["version"]
    await client.put(url, json={"expected_version": v, "tds": "700", "note": "keep me"}, headers=payroll["hr"])
    # the manager becomes eligible again, the employee leaves payroll
    r = await client.put(
        f"/api/admin/payroll/employees/{uid['manager']}/details",
        json={"include_in_payroll": True, "current_password": PASSWORD},
        headers=payroll["admin"],
    )
    assert r.status_code == 200
    await exclude_from_payroll(client, payroll["admin"], uid["employee"])
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/refresh", headers=payroll["hr"])
    assert r.status_code == 200 and r.json()["counts"]["employees"] == 2
    items = await payslips_of(client, payroll["hr"], run["id"])
    emp = next(i for i in items if i["user"]["employee_id"] == "EMP001")
    assert "NOT_ELIGIBLE" in emp["warnings"]
    got = (await client.get(url, headers=payroll["hr"])).json()
    assert got["inputs"]["tds"] == "700.00" and got["inputs"]["note"] == "keep me"


async def test_reopen_returns_a_processed_run_to_draft(client, payroll):
    run = await new_run(client, payroll["hr"])
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    assert r.json()["status"] == "processed" and r.json()["counts"]["processed"] == 2
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    r = await client.put(
        f"/api/admin/payroll/payslips/{row['id']}", json={"expected_version": 1, "note": "x"}, headers=payroll["hr"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"  # not a draft any more
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/reopen", headers=payroll["hr"])
    assert r.status_code == 200 and r.json()["status"] == "draft" and r.json()["counts"]["draft"] == 2
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/reopen", headers=payroll["hr"])
    assert r.status_code == 409  # not processed


async def test_delete_draft_run_and_payslip(client, payroll, db):
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    assert (await client.delete(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).status_code == 204
    assert (await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).status_code == 404
    assert (await client.delete(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])).status_code == 204
    assert (await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])).status_code == 404
    assert (await db.execute(select(PayrollRun))).first() is None


# ================================================================================================ publish
async def test_publish_requires_processed_run_password_and_locks(client, payroll, db):
    run = await new_run(client, payroll["hr"])
    r = await publish_run(client, payroll["hr"], run["id"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"  # not processed yet
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/publish", json={}, headers=payroll["hr"])
    assert r.status_code == 422
    r = await publish_run(client, payroll["hr"], run["id"], pw="Wrong-password-9!")
    assert r.status_code == 403 and r.json()["error"]["code"] == "REAUTH_FAILED"
    assert (await client.get(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])).json()["status"] == "processed"
    notifications.OUTBOX.clear()
    r = await publish_run(client, payroll["hr"], run["id"])
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["status"] == "published" and out["counts"]["published"] == 2 and out["company_name"] is None
    assert (
        out["published_by"]["employee_id"] == "HR001"
        and out["totals"]["net_pay"] == "57800.00"
        and out["totals"]["employer_pf"] == "3600.00"
    )
    # idempotency / locking
    r = await publish_run(client, payroll["hr"], run["id"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    for path in ("refresh", "process", "reopen"):
        r = await client.post(f"/api/admin/payroll/runs/{run['id']}/{path}", headers=payroll["hr"])
        assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED", path
    assert (await client.delete(f"/api/admin/payroll/runs/{run['id']}", headers=payroll["hr"])).status_code == 409
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    r = await client.put(
        f"/api/admin/payroll/payslips/{row['id']}", json={"expected_version": 1, "note": "x"}, headers=payroll["hr"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    assert (await client.delete(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).status_code == 409
    # push: data only, no amounts
    pushes = [e for e in notifications.OUTBOX if e["payload"]["type"] == "payslip_published"]
    assert len(pushes) == 2
    p = pushes[0]["payload"]
    assert (
        set(p) == {"type", "id", "title", "body"}
        and p["title"] == "Payslip available"
        and p["body"] == "Your payslip for June 2025 is ready."
    )


async def test_nothing_to_publish_and_run_without_company_name(client, payroll, uid):
    await exclude_from_payroll(client, payroll["admin"], uid["employee"])
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    run = await new_run(client, payroll["hr"])
    assert run["counts"]["employees"] == 0
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    r = await publish_run(client, payroll["hr"], run["id"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "NOTHING_TO_PUBLISH"


async def test_publish_snapshots_company_name(client, payroll, db):
    from app.models import Company

    c = await db.get(Company, 1)
    c.name = "Aksan Capital Advisory Private Limited"
    await db.commit()
    run = await new_run(client, payroll["hr"])
    out = await process_and_publish(client, payroll["hr"], run["id"])
    assert out["company_name"] == "Aksan Capital Advisory Private Limited"
    c.name = "Renamed Ltd"
    await db.commit()
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    p = (await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])).json()
    assert p["company"]["name"] == "Aksan Capital Advisory Private Limited"


async def test_maker_checker_second_person(client, payroll):
    r = await client.put(  # changing the rule needs role.assign (admin) and the caller's own password
        "/api/admin/payroll/settings",
        json={"settings": {"payroll.publish_requires_second_person": True}, "current_password": PASSWORD},
        headers=payroll["admin"],
    )
    assert r.status_code == 200, r.text
    run = await new_run(client, payroll["hr"])
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    r = await publish_run(client, payroll["hr"], run["id"])
    assert r.status_code == 403 and r.json()["error"]["code"] == "SECOND_PERSON_REQUIRED"
    r = await publish_run(client, payroll["admin"], run["id"])
    assert r.status_code == 200


# ================================================================================================ employee side
async def published_run(client, payroll):
    run = await new_run(client, payroll["hr"])
    await process_and_publish(client, payroll["hr"], run["id"])
    return run


async def test_employee_list_has_no_amounts_and_only_own_published(client, payroll, uid):
    run = await new_run(client, payroll["hr"])
    r = await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])
    assert r.status_code == 200 and r.json() == {"year": 2025, "items": []}  # drafts are invisible
    await process_and_publish(client, payroll["hr"], run["id"])
    r = await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])
    body = r.json()
    assert len(body["items"]) == 1
    item = body["items"][0]
    assert set(item) == {"id", "period", "revision", "published_at", "revised", "pay_date"}
    assert item["period"] == PERIOD and item["revised"] is False and item["pay_date"] == "2025-06-30"
    assert not any(k in r.text for k in ("net", "gross", "amount", "23200"))
    assert (await client.get("/api/payslips", params={"year": 2024}, headers=payroll["employee"])).json()["items"] == []
    assert (await client.get("/api/payslips", params={"year": 1999}, headers=payroll["employee"])).status_code == 422


async def test_employee_detail_shape_masking_and_audit(client, payroll, uid, db):
    await client.put(f"/api/admin/payroll/employees/{uid['employee']}/details",
                     json={"designation": "Analyst", "pan": "ABCDE1234F", "uan": "100200300400", "bank_name": "HDFC Bank",
                           "bank_account": "123456789012", "bank_ifsc": "HDFC0001234", "current_password": PASSWORD}, headers=payroll["hr"])  # fmt: skip
    run = await published_run(client, payroll)
    item = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"][0]
    r = await client.get(f"/api/payslips/{item['id']}", headers=payroll["employee"])
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["status"] == "published" and p["period"] == PERIOD and p["revision"] == 1 and p["revision_info"] is None
    assert (
        p["employee"]["pan"] == "XXXXX1234X"
        and p["employee"]["bank_account"] == "XXXXXXXX9012"
        and p["employee"]["designation"] == "Analyst"
    )
    assert p["employee"]["uan"] == "100200300400" and p["employee"]["employee_id"] == "EMP001"
    assert p["days"] == {"basis_days": 30.0, "employed_days": 30.0, "lop_days": 0.0, "paid_days": 30.0}
    assert [e["code"] for e in p["earnings"]] == ["BASIC", "HRA"] and p["earnings"][0] == {
        "code": "BASIC",
        "label": "Basic",
        "monthly": "20000.00",
        "amount": "20000.00",
    }
    assert p["deductions"] == [
        {"code": "PF", "label": "PF - employee share", "amount": "1800.00"},
        {"code": "PF_ER_REC", "label": "PF - employer share (recovered)", "amount": "1800.00"},
    ]
    assert p["employer_contributions"] is None  # show_employer_contributions is off
    assert p["totals"] == {"gross_earnings": "25000.00", "total_deductions": "3600.00", "net_pay": "21400.00",
                           "net_pay_words": "Rupees Twenty-One Thousand Four Hundred Only"}  # fmt: skip
    assert "errors" not in p and "inputs" not in p and "version" not in p and "lop" not in p  # admin-only fields
    assert any(a.action == "payroll.payslip.view_own" for a in (await db.execute(select(AuditLog))).scalars())
    assert run["id"]


async def test_pan_unmasked_when_setting_off_and_employer_contributions_when_on(client, payroll, uid):
    await set_setting(client, payroll["hr"], mask_pan=False, show_employer_contributions=True)
    await client.put(
        f"/api/admin/payroll/employees/{uid['employee']}/details",
        json={"pan": "ABCDE1234F", "current_password": PASSWORD},
        headers=payroll["hr"],
    )
    await published_run(client, payroll)
    item = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"][0]
    p = (await client.get(f"/api/payslips/{item['id']}", headers=payroll["employee"])).json()
    assert p["employee"]["pan"] == "ABCDE1234F"
    assert p["employer_contributions"] == [{"code": "PF_ER", "label": "Employer PF", "amount": "1800.00"}]


async def test_employee_cannot_see_others_drafts_or_unknown_payslips(client, payroll, uid, db):
    run = await new_run(client, payroll["hr"])
    emp_row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    mgr_row = await row_for(client, payroll["hr"], run["id"], "MGR001")
    # own draft: 404 (not published)
    for path in (f"/api/payslips/{emp_row['id']}", f"/api/payslips/{emp_row['id']}/pdf"):
        assert (await client.get(path, headers=payroll["employee"])).status_code == 404
    await process_and_publish(client, payroll["hr"], run["id"])
    # someone else's published payslip: 404, same as unknown (no existence leak)
    for path in (f"/api/payslips/{mgr_row['id']}", f"/api/payslips/{mgr_row['id']}/pdf"):
        r = await client.get(path, headers=payroll["employee"])
        assert r.status_code == 404 and r.json()["error"]["code"] == "NOT_FOUND"
    r = await client.get("/api/payslips/00000000-0000-0000-0000-000000000000", headers=payroll["employee"])
    assert r.status_code == 404
    # HR does not get other people's payslips through the employee routes either
    assert (await client.get(f"/api/payslips/{emp_row['id']}", headers=payroll["hr"])).status_code == 404
    assert (await client.get(f"/api/payslips/{emp_row['id']}/pdf", headers=payroll["hr"])).status_code == 404
    denied = (await db.execute(select(AuditLog).where(AuditLog.action == "payroll.denied"))).scalars().all()
    assert denied and all(d.outcome == "failure" for d in denied)
    # the manager sees only his own
    mine = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["manager"])).json()["items"]
    assert [i["id"] for i in mine] == [mgr_row["id"]]


async def test_employee_pdf_headers_and_content(client, payroll, db):
    await published_run(client, payroll)
    item = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"][0]
    r = await client.get(f"/api/payslips/{item['id']}/pdf", headers=payroll["employee"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == 'attachment; filename="payslip-EMP001-2025-06.pdf"'
    assert r.headers["cache-control"] == "no-store" and r.headers["x-content-type-options"] == "nosniff"
    assert (
        r.content.startswith(b"%PDF-")
        and b"/JavaScript" not in r.content
        and b"/URI" not in r.content
        and b"/EmbeddedFile" not in r.content
    )
    assert b"Payslip 2025-06" in r.content  # the only metadata title
    assert any(a.action == "payroll.payslip.download" for a in (await db.execute(select(AuditLog))).scalars())
    # drafts of others / the same endpoint for HR's admin copy
    row = await row_for(
        client,
        payroll["hr"],
        (await client.get("/api/admin/payroll/runs", params={"year": 2025}, headers=payroll["hr"])).json()["items"][0]["id"],
        "EMP001",
    )
    r = await client.get(f"/api/admin/payroll/payslips/{row['id']}/pdf", headers=payroll["hr"])
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store" and r.content.startswith(b"%PDF-")


async def test_employee_detail_rate_limit(client, payroll):
    await published_run(client, payroll)
    item = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"][0]
    codes = [(await client.get(f"/api/payslips/{item['id']}", headers=payroll["employee"])).status_code for _ in range(32)]
    assert codes[:30] == [200] * 30 and codes[30:] == [429, 429]


# ================================================================================================ revisions
async def test_revision_flow_supersedes_in_one_step(client, payroll, uid, db):
    run = await published_run(client, payroll)
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    old_id = row["id"]
    notifications.OUTBOX.clear()
    # revise: reason is required (3-200)
    url = f"/api/admin/payroll/runs/{run['id']}/employees/{uid['employee']}/revise"
    assert (await client.post(url, json={"reason": "ab"}, headers=payroll["hr"])).status_code == 422
    r = await client.post(url, json={"reason": "Leave approved after payroll"}, headers=payroll["hr"])
    assert r.status_code == 201, r.text
    rev = r.json()
    assert rev["revision"] == 2 and rev["status"] == "draft" and rev["revision_info"]["reason"] == "Leave approved after payroll"
    assert rev["revision_info"]["supersedes_id"] == old_id
    # a second open revision is refused
    r = await client.post(url, json={"reason": "again please"}, headers=payroll["hr"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "CONFLICT"
    # while the revision is open the employee still sees revision 1 only
    mine = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"]
    assert [i["revision"] for i in mine] == [1]
    # edit, process, publish the single revision
    r = await client.put(
        f"/api/admin/payroll/payslips/{rev['id']}",
        json={"expected_version": rev["version"], "adjustments": [{"code": "ARREARS", "label": "Back pay", "amount": "1000"}]},
        headers=payroll["hr"],
    )
    assert r.status_code == 200 and r.json()["totals"]["gross_earnings"] == "26000.00"
    r = await client.post(f"/api/admin/payroll/payslips/{rev['id']}/process", headers=payroll["hr"])
    assert r.status_code == 200 and r.json()["status"] == "processed"
    r = await client.post(f"/api/admin/payroll/payslips/{rev['id']}/publish", json={}, headers=payroll["hr"])
    assert r.status_code == 422
    r = await client.post(
        f"/api/admin/payroll/payslips/{rev['id']}/publish", json={"current_password": "Wrong-password-9!"}, headers=payroll["hr"]
    )
    assert r.status_code == 403
    r = await client.post(
        f"/api/admin/payroll/payslips/{rev['id']}/publish", json={"current_password": PASSWORD}, headers=payroll["hr"]
    )
    assert r.status_code == 200 and r.json()["status"] == "published"
    old = (await client.get(f"/api/admin/payroll/payslips/{old_id}", headers=payroll["hr"])).json()
    assert old["status"] == "superseded" and old["revision_info"]["superseded_by_id"] == rev["id"]
    # employee: only the live revision in the list, the old one is still readable
    mine = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"]
    assert [(i["id"], i["revision"], i["revised"]) for i in mine] == [(rev["id"], 2, True)]
    r = await client.get(f"/api/payslips/{old_id}", headers=payroll["employee"])
    assert (
        r.status_code == 200 and r.json()["status"] == "superseded" and r.json()["revision_info"]["superseded_by_id"] == rev["id"]
    )
    new = (await client.get(f"/api/payslips/{rev['id']}", headers=payroll["employee"])).json()
    assert new["totals"]["gross_earnings"] == "26000.00" and new["revision_info"]["reason"] == "Leave approved after payroll"
    pdf = await client.get(f"/api/payslips/{rev['id']}/pdf", headers=payroll["employee"])
    assert 'filename="payslip-EMP001-2025-06-r2.pdf"' in pdf.headers["content-disposition"]
    push = [e["payload"] for e in notifications.OUTBOX if e["payload"]["type"] == "payslip_published"]
    assert push == [
        {
            "type": "payslip_published",
            "id": rev["id"],
            "title": "Revised payslip available",
            "body": "Your payslip for June 2025 is ready.",
        }
    ]
    # the published revision is locked like any published payslip
    r = await client.put(
        f"/api/admin/payroll/payslips/{rev['id']}", json={"expected_version": 1, "note": "x"}, headers=payroll["hr"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    # exactly one live copy of the employee's payslip in the run
    copies = (
        (
            await db.execute(
                select(Payslip.status)
                .where(Payslip.user_id == __import__("uuid").UUID(uid["employee"]))
                .order_by(Payslip.revision)
            )
        )
        .scalars()
        .all()
    )
    assert copies == ["superseded", "published"]


async def test_revise_requires_published_run_and_creates_missing_employee_revision_one(client, payroll, uid):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    run = await new_run(client, payroll["hr"])
    url = f"/api/admin/payroll/runs/{run['id']}/employees/{uid['employee']}/revise"
    r = await client.post(url, json={"reason": "not yet published"}, headers=payroll["hr"])
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"
    await process_and_publish(client, payroll["hr"], run["id"])
    # the manager was not in the run; HR adds him later as revision 1 (a late joiner to the run)
    r = await client.put(
        f"/api/admin/payroll/employees/{uid['manager']}/details",
        json={"include_in_payroll": True, "current_password": PASSWORD},
        headers=payroll["admin"],
    )
    assert r.status_code == 200
    r = await client.post(
        f"/api/admin/payroll/runs/{run['id']}/employees/{uid['manager']}/revise",
        json={"reason": "Missed in the run"},
        headers=payroll["hr"],
    )
    assert r.status_code == 201 and r.json()["revision"] == 1 and r.json()["status"] == "draft"
    assert r.json()["revision_info"] is None


async def test_held_payslip_is_published_on_its_own_after_the_run(client, payroll):
    run = await new_run(client, payroll["hr"])
    mgr = await row_for(client, payroll["hr"], run["id"], "MGR001")
    v = (await client.get(f"/api/admin/payroll/payslips/{mgr['id']}", headers=payroll["hr"])).json()["version"]
    await client.put(
        f"/api/admin/payroll/payslips/{mgr['id']}", json={"expected_version": v, "hold": True}, headers=payroll["hr"]
    )
    out = await process_and_publish(client, payroll["hr"], run["id"])
    assert out["counts"]["published"] == 1 and out["counts"]["held"] == 1
    assert (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["manager"])).json()["items"] == []
    r = await client.post(f"/api/admin/payroll/payslips/{mgr['id']}/process", headers=payroll["hr"])
    assert r.status_code == 200 and r.json()["status"] == "processed" and r.json()["hold"] is False
    r = await client.post(f"/api/admin/payroll/payslips/{mgr['id']}/reopen", headers=payroll["hr"])
    assert r.json()["status"] == "draft"
    await client.post(f"/api/admin/payroll/payslips/{mgr['id']}/process", headers=payroll["hr"])
    r = await client.post(
        f"/api/admin/payroll/payslips/{mgr['id']}/publish", json={"current_password": PASSWORD}, headers=payroll["hr"]
    )
    assert r.status_code == 200 and r.json()["status"] == "published"
    assert len((await client.get("/api/payslips", params={"year": 2025}, headers=payroll["manager"])).json()["items"]) == 1
    r = await client.post(
        f"/api/admin/payroll/payslips/{mgr['id']}/publish", json={"current_password": PASSWORD}, headers=payroll["hr"]
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "PAYROLL_LOCKED"


# ================================================================================================ self rule
async def test_self_rule_on_payslips(client, payroll, uid):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    await client.put(
        f"/api/admin/payroll/employees/{uid['hr']}/details",
        json={"include_in_payroll": True, "current_password": PASSWORD},
        headers=payroll["admin"],
    )
    await add_structure(client, payroll["admin"], uid["hr"], basic="40000")
    run = await new_run(client, payroll["hr"])
    own = await row_for(client, payroll["hr"], run["id"], "HR001")
    url = f"/api/admin/payroll/payslips/{own['id']}"
    for method, path, body in (
        ("PUT", url, {"expected_version": 1, "note": "x"}),
        ("DELETE", url, None),
    ):  # fmt: skip
        r = await client.request(method, path, json=body, headers=payroll["hr"])
        assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN", method
    # the admin (another person) can edit HR's payslip; HR can still view it
    assert (await client.get(url, headers=payroll["hr"])).status_code == 200
    v = (await client.get(url, headers=payroll["admin"])).json()["version"]
    assert (await client.put(url, json={"expected_version": v, "note": "ok"}, headers=payroll["admin"])).status_code == 200
    # run level publish may include the caller's own payslip (audited self_included)
    await client.post(f"/api/admin/payroll/runs/{run['id']}/process", headers=payroll["hr"])
    notifications.OUTBOX.clear()
    out = await publish_run(client, payroll["hr"], run["id"])
    assert out.status_code == 200
    # single payslip actions on your own payslip stay forbidden, revise too
    r = await client.post(
        f"/api/admin/payroll/runs/{run['id']}/employees/{uid['hr']}/revise", json={"reason": "my own pay"}, headers=payroll["hr"]
    )
    assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN"
    for action in ("process", "reopen"):
        assert (await client.post(f"{url}/{action}", headers=payroll["hr"])).status_code == 403
    assert (await client.post(f"{url}/publish", json={"current_password": PASSWORD}, headers=payroll["hr"])).status_code == 403
    r = await client.post(
        f"/api/admin/payroll/runs/{run['id']}/employees/{uid['hr']}/revise",
        json={"reason": "by the admin"},
        headers=payroll["admin"],
    )
    assert r.status_code == 201


async def test_self_included_is_audited_on_run_publish(client, payroll, uid, db):
    await exclude_from_payroll(client, payroll["admin"], uid["manager"])
    await client.put(
        f"/api/admin/payroll/employees/{uid['hr']}/details",
        json={"include_in_payroll": True, "current_password": PASSWORD},
        headers=payroll["admin"],
    )
    await add_structure(client, payroll["admin"], uid["hr"], basic="40000")
    run = await new_run(client, payroll["hr"])
    await process_and_publish(client, payroll["hr"], run["id"])
    row = (await db.execute(select(AuditLog).where(AuditLog.action == "payroll.run.publish"))).scalar_one()
    assert row.after["self_included"] is True and row.after["payslips"] == 2


# ================================================================================================ audit / encryption
async def test_audit_rows_never_hold_amounts_bank_or_pan(client, payroll, uid, db):
    await client.put(
        f"/api/admin/payroll/employees/{uid['employee']}/details",
        json={
            "pan": "ABCDE1234F",
            "bank_account": "123456789012",
            "bank_ifsc": "HDFC0001234",
            "bank_name": "HDFC Bank",
            "current_password": PASSWORD,
        },
        headers=payroll["hr"],
    )
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    await client.put(
        f"/api/admin/payroll/payslips/{row['id']}",
        json={"expected_version": 1, "tds": "777", "adjustments": [{"code": "BONUS", "label": "B", "amount": "4242"}]},
        headers=payroll["hr"],
    )
    await client.get(f"/api/admin/payroll/payslips/{row['id']}", headers=payroll["hr"])
    await process_and_publish(client, payroll["hr"], run["id"])
    item = (await client.get("/api/payslips", params={"year": 2025}, headers=payroll["employee"])).json()["items"][0]
    await client.get(f"/api/payslips/{item['id']}", headers=payroll["employee"])
    await client.get(f"/api/payslips/{item['id']}/pdf", headers=payroll["employee"])
    logs = (await db.execute(select(AuditLog))).scalars().all()
    blob = " ".join(f"{x.action} {x.before} {x.after}" for x in logs)
    for needle in ("ABCDE1234F", "123456789012", "HDFC", "777", "4242", "23200", "25000", "20000", "1800"):
        assert needle not in blob, needle
    actions = {x.action for x in logs}
    required = {"payroll.run.create", "payroll.payslip.update", "payroll.payslip.view", "payroll.run.process", "payroll.run.publish",
                "payroll.payslip.view_own", "payroll.payslip.download", "payroll.details.update", "payroll.settings.update"}  # fmt: skip
    assert required <= actions, required - actions


async def test_money_and_identity_are_encrypted_at_rest(client, payroll, uid, db):
    await client.put(
        f"/api/admin/payroll/employees/{uid['employee']}/details",
        json={"pan": "ABCDE1234F", "bank_account": "123456789012", "current_password": PASSWORD},
        headers=payroll["hr"],
    )
    run = await new_run(client, payroll["hr"])
    raw = (await db.execute(text("SELECT result_enc, inputs_enc, identity_enc FROM payslips"))).all()
    assert raw
    for result, inputs, identity in raw:
        assert b"25000" not in bytes(result) and b"gross" not in bytes(result)
        assert b"tds" not in bytes(inputs)
        assert identity is None or b"ABCDE1234F" not in bytes(identity)
    comp = (await db.execute(text("SELECT components_enc FROM salary_structures"))).scalars().all()
    assert comp and all(b"20000" not in bytes(c) and b"basic" not in bytes(c) for c in comp)
    assert run["id"]


async def test_revision_reason_and_lop_reason_encrypted(client, payroll, uid, db):
    run = await published_run(client, payroll)
    await client.post(
        f"/api/admin/payroll/runs/{run['id']}/employees/{uid['employee']}/revise",
        json={"reason": "Secret reason text"},
        headers=payroll["hr"],
    )
    raw = (
        (await db.execute(text("SELECT revision_reason_enc FROM payslips WHERE revision_reason_enc IS NOT NULL"))).scalars().all()
    )
    assert raw and all(b"Secret" not in bytes(x) for x in raw)


async def test_details_row_created_lazily(client, payroll, uid, db):
    assert (await db.get(EmployeePayrollDetails, __import__("uuid").UUID(uid["employee"]))) is None
    await new_run(client, payroll["hr"])
    assert (await db.get(EmployeePayrollDetails, __import__("uuid").UUID(uid["employee"]))) is None  # reading creates nothing


# ================================================================================================ register
async def test_register_workbook_sheets_totals_and_headers(client, payroll, uid):
    import io

    from openpyxl import load_workbook

    await add_structure(client, payroll["hr"], uid["employee"], basic="20000", hra="5000", eff="2025-02-01",
                        allowances=[{"code": "SPECIAL", "name": "Special", "amount": "3000", "prorate": True, "esi_wage": True}],
                        fixed_deductions=[{"code": "CANTEEN", "name": "Canteen", "amount": "500"}])  # fmt: skip
    run = await new_run(client, payroll["hr"])
    row = await row_for(client, payroll["hr"], run["id"], "EMP001")
    await client.put(
        f"/api/admin/payroll/payslips/{row['id']}",
        json={"expected_version": 1, "adjustments": [{"code": "BONUS", "label": '=HYPERLINK("x")', "amount": "1000"}]},
        headers=payroll["hr"],
    )
    await process_and_publish(client, payroll["hr"], run["id"])
    r = await client.post(f"/api/admin/payroll/runs/{run['id']}/register", json={}, headers=payroll["hr"])
    assert r.status_code == 200, r.text
    assert (
        r.headers["content-disposition"] == 'attachment; filename="payroll-register-2025-06.xlsx"'
        and r.headers["cache-control"] == "no-store"
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Register", "Summary"]
    ws = wb["Register"]
    head = [c.value for c in ws[1]]
    assert head[:7] == ["Employee ID", "Name", "Department", "Designation", "Status", "Revision", "Hold"]
    for col in (
        "BASIC",
        "HRA",
        "SPECIAL",
        "BONUS",
        "Gross",
        "PF",
        "ESI",
        "PT",
        "TDS",
        "CANTEEN",
        "Total Deductions",
        "Net Pay",
        "Employer PF",
        "Employer ESI",
    ):
        assert col in head, col
    rows = {r_[0].value: {h: c.value for h, c in zip(head, r_, strict=True)} for r_ in ws.iter_rows(min_row=2)}
    e = rows["EMP001"]
    assert (
        e["Gross"] == 29000 and e["BONUS"] == 1000 and e["CANTEEN"] == 500 and e["Status"] == "published"
    )  # 20000+5000+3000+1000
    assert isinstance(e["Net Pay"], (int, float)) and e["Net Pay"] == e["Gross"] - e["Total Deductions"]
    summary = {a.value: b.value for a, b in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Employees"] == 2 and summary["Total Gross"] == e["Gross"] + rows["MGR001"]["Gross"]
    assert "Bank" not in wb.sheetnames


async def test_register_bank_sheet_needs_manage_and_password(client, payroll, uid, db):
    import io

    from openpyxl import load_workbook

    from app.core import security
    from app.models import Permission, Role, RolePermission, User

    await client.put(
        f"/api/admin/payroll/employees/{uid['employee']}/details",
        json={"bank_name": "HDFC Bank", "bank_account": "123456789012", "bank_ifsc": "HDFC0001234", "current_password": PASSWORD},
        headers=payroll["hr"],
    )
    run = await new_run(client, payroll["hr"])
    await process_and_publish(client, payroll["hr"], run["id"])
    url = f"/api/admin/payroll/runs/{run['id']}/register"
    r = await client.post(url, json={"include_bank": True}, headers=payroll["hr"])
    assert r.status_code == 422 and r.json()["error"]["details"][0]["field"] == "current_password"
    r = await client.post(url, json={"include_bank": True, "current_password": "Wrong-password-9!"}, headers=payroll["hr"])
    assert r.status_code == 403 and r.json()["error"]["code"] == "REAUTH_FAILED"
    r = await client.post(url, json={"include_bank": True, "current_password": PASSWORD}, headers=payroll["hr"])
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Register", "Summary", "Bank"]
    bank = {row[0].value: [c.value for c in row] for row in wb["Bank"].iter_rows(min_row=2)}
    assert (
        bank["EMP001"][:5] == ["EMP001", "Test Employee", "HDFC Bank", "123456789012", "HDFC0001234"]
        and bank["EMP001"][6] == "bank_transfer"
    )
    # a read-only viewer (payroll.view_all, no payroll.manage) gets the register but never the bank sheet
    role = Role(name="auditor", description="read only payroll", is_system=False)
    db.add(role)
    await db.flush()
    perm = (await db.execute(select(Permission).where(Permission.code == "payroll.view_all"))).scalar_one()
    db.add(RolePermission(role_id=role.id, permission_id=perm.id))
    db.add(
        User(
            employee_id="AUD001",
            password_hash=security.hash_password(PASSWORD),
            full_name="Aud Itor",
            role_id=role.id,
            must_change_password=False,
            status="active",
        )
    )
    await db.commit()
    tok = (await client.post("/api/auth/login", json={"employee_id": "AUD001", "password": PASSWORD})).json()
    aud = bearer(tok)
    assert (await client.post(url, json={}, headers=aud)).status_code == 200
    r = await client.post(url, json={"include_bank": True, "current_password": PASSWORD}, headers=aud)
    assert r.status_code == 403 and r.json()["error"]["code"] == "FORBIDDEN"
    assert (await client.get("/api/admin/payroll/runs", headers=aud)).status_code == 200
    assert (await client.post("/api/admin/payroll/runs", json={"period": "2025-05"}, headers=aud)).status_code == 403
    assert (
        await client.post(f"/api/admin/payroll/runs/{run['id']}/publish", json={"current_password": PASSWORD}, headers=aud)
    ).status_code == 403
    logs = (await db.execute(select(AuditLog).where(AuditLog.action == "payroll.register.export"))).scalars().all()
    assert sorted(x.after["include_bank"] for x in logs) == [False, True]
    assert all("123456789012" not in str(x.after) for x in logs)


async def test_register_rate_limit_is_5_per_minute(client, payroll):
    run = await new_run(client, payroll["hr"])
    codes = [
        (await client.post(f"/api/admin/payroll/runs/{run['id']}/register", json={}, headers=payroll["hr"])).status_code
        for _ in range(6)
    ]
    assert codes[:5] == [200] * 5 and codes[5] == 429


def test_license_read_only_mode_treats_register_as_a_read():
    from app.services import license

    assert any(m == "POST" and rx.match("/api/admin/payroll/runs/abc/register") for m, rx in license._READ_ONLY_EXEMPT)
    assert not any(rx.match("/api/admin/payroll/runs/abc/publish") for _, rx in license._READ_ONLY_EXEMPT)


# ================================================================================================ misc
async def test_joiner_and_leaver_are_paid_pro_rata(client, payroll, uid, db):
    from app.models import User

    emp = await db.get(User, __import__("uuid").UUID(uid["employee"]))
    emp.join_date = dt.date(2025, 6, 11)  # 20 of 30 days
    mgr = await db.get(User, __import__("uuid").UUID(uid["manager"]))
    mgr.status, mgr.deactivated_at, mgr.deactivation_reason = (
        "inactive",
        dt.datetime(2025, 6, 16, 5, 0, tzinfo=dt.UTC),
        "resigned",
    )  # last day 15th
    await db.commit()
    run = await new_run(client, payroll["hr"])
    items = {i["user"]["employee_id"]: i for i in await payslips_of(client, payroll["hr"], run["id"])}
    assert items["EMP001"]["paid_days"] == 20.0
    p = (await client.get(f"/api/admin/payroll/payslips/{items['EMP001']['id']}", headers=payroll["hr"])).json()
    assert p["lop"]["breakdown"]["not_employed"] == 10.0 and p["earnings"][0]["amount"] == "13333.00"  # 20000 x 20/30
    assert items["MGR001"]["paid_days"] == 15.0
    # a leaver deactivated before the month starts is not in the run at all
    mgr.deactivated_at = dt.datetime(2025, 5, 20, 5, 0, tzinfo=dt.UTC)
    await db.commit()
    run2 = await new_run(client, payroll["hr"], period="2025-07")
    assert {i["user"]["employee_id"] for i in await payslips_of(client, payroll["hr"], run2["id"])} == {"EMP001"}


async def test_run_list_shape(client, payroll):
    run = await new_run(client, payroll["hr"])
    r = await client.get("/api/admin/payroll/runs", params={"year": 2025}, headers=payroll["hr"])
    assert r.json()["items"] == [
        {
            "id": run["id"],
            "period": PERIOD,
            "status": "draft",
            "employees": 2,
            "net_pay_total": run["totals"]["net_pay"],
            "published_at": None,
        }
    ]
    assert run["counts"].keys() == {
        "employees",
        "draft",
        "processed",
        "published",
        "held",
        "errors",
        "warnings",
        "open_revisions",
    }
    assert run["totals"].keys() == {
        "gross_earnings",
        "total_deductions",
        "net_pay",
        "employer_pf",
        "employer_esi",
        "employer_recovered",
        "employer_cost",
    }


async def test_payslip_list_filters(client, payroll, uid):
    run = await new_run(client, payroll["hr"])
    assert len(await payslips_of(client, payroll["hr"], run["id"], q="manager")) == 1
    assert len(await payslips_of(client, payroll["hr"], run["id"], flag="warning")) == 2
    assert await payslips_of(client, payroll["hr"], run["id"], flag="error") == []
    assert await payslips_of(client, payroll["hr"], run["id"], flag="hold") == []
    assert await payslips_of(client, payroll["hr"], run["id"], status="processed") == []
    r = await client.get(f"/api/admin/payroll/runs/{run['id']}/payslips", params={"flag": "bogus"}, headers=payroll["hr"])
    assert r.status_code == 422
