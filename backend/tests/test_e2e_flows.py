"""Cross-feature end-to-end flows (final QA pass). Each test drives several routers through the public API only.

Clock: 2026-10-21 (Wednesday), 10:00 IST = 04:30Z. Shift 09:30-18:30, grace 10 -> a 10:00 punch is 30 min late.
"""

from __future__ import annotations

import datetime as dt
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.clock import utcnow as real_utcnow
from app.core.config import get_settings
from app.models import AuditLog, LeaveType
from app.services import notifications
from tests.conftest import USERS, bearer, do_login

NOW = dt.datetime(2026, 10, 21, 4, 30, tzinfo=dt.UTC)
NEAR = {"lat": 12.9716 + 0.0009, "lng": 77.5946, "accuracy_m": 10}


@pytest.fixture
def clock(monkeypatch):
    state = {"now": NOW}
    for mod in ("attendance", "late", "leave", "food"):
        monkeypatch.setattr(f"app.services.{mod}.utcnow", lambda: state["now"])
    st = get_settings()
    monkeypatch.setattr(st, "leave_write_rate_per_min", 1000)
    monkeypatch.setattr(st, "leave_write_rate_per_day", 100000)
    monkeypatch.setattr(st, "leave_decide_rate_per_min", 1000)
    notifications.OUTBOX.clear()
    return state


@pytest.fixture
async def H(tokens_for):
    return {r: bearer(await tokens_for(r)) for r in USERS}


@pytest.fixture(autouse=True)
def storage(tmp_path, monkeypatch):
    root = tmp_path / "store"
    root.mkdir()
    monkeypatch.setattr(get_settings(), "file_storage_dir", str(root))


def sheet_rows(content: bytes, name: str) -> list[dict]:
    wb = load_workbook(io.BytesIO(content))
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r, strict=True)) for r in rows[1:]]


# ---------------------------------------------------------------- login -> punch -> late -> leave -> approval -> export
async def test_workday_flow_login_punch_late_leave_approval_export(client, H, clock, db):
    # login (real endpoint) and identity
    r = await do_login(client, USERS["employee"], device_id="dev-1", device_name="Pixel")
    assert r.status_code == 200
    emp = bearer(r.json())
    me = (await client.get("/api/auth/me", headers=emp)).json()
    assert me["employee_id"] == "EMP001" and me["role"] == "employee"

    # punch in 30 minutes after shift start -> late
    r = await client.post("/api/attendance/check-in", headers=emp, json={**NEAR, "fix_time": NOW.isoformat()})
    assert r.status_code == 201, r.text
    rec = r.json()
    assert rec["status"] == "late" and rec["minutes_late"] == 30 and rec["work_date"] == "2026-10-21"
    assert (await client.post("/api/attendance/check-in", headers=emp, json={**NEAR, "fix_time": NOW.isoformat()})).json()[
        "error"
    ]["code"] == "ALREADY_CHECKED_IN"

    # the late feed, summary, and the manager's team view all agree
    late = (await client.get("/api/late?from=2026-10-01&to=2026-10-31", headers=emp)).json()
    assert late["total"] == 1 and late["items"][0]["minutes_late"] == 30 and late["items"][0]["grace_minutes"] == 10
    summ = (await client.get("/api/late/summary?month=2026-10", headers=emp)).json()
    assert summ["late_count"] == 1 and summ["total_minutes_late"] == 30
    team = (await client.get("/api/late/team?from=2026-10-01&to=2026-10-31", headers=H["manager"])).json()
    assert [i["user"]["employee_id"] for i in team["items"]] == ["EMP001"]
    assert (await client.get("/api/late/team", headers=emp)).status_code == 403

    # check out (worked 9 h) keeps the late status
    clock["now"] = NOW + dt.timedelta(hours=9)
    r = await client.post("/api/attendance/check-out", headers=emp, json={**NEAR, "fix_time": clock["now"].isoformat()})
    assert r.status_code == 200 and r.json()["status"] == "late" and r.json()["worked_minutes"] == 540
    clock["now"] = NOW

    # leave: apply Mon 26 - Tue 27 Oct (2 working days)
    lt = LeaveType(code="CL", name="Casual Leave", annual_quota=12, is_paid=True, allow_half_day=True)
    db.add(lt)
    await db.commit()
    r = await client.post(
        "/api/leave/requests",
        headers=emp,
        json={"leave_type_id": str(lt.id), "from_date": "2026-10-26", "to_date": "2026-10-27", "reason": "vacation"},
    )
    assert r.status_code == 201, r.text
    req = r.json()
    assert req["days"] == 2.0 and req["status"] == "pending"
    bal = (await client.get("/api/leave/balance", headers=emp)).json()["leave"][0]
    assert (bal["used"], bal["pending"], bal["remaining"]) == (0.0, 2.0, 10.0)
    # approval pushes go to the manager; the employee cannot approve (RBAC), nor can they decide their own request
    assert any(o["payload"]["type"] == "new_request" for o in notifications.OUTBOX)
    url = f"/api/admin/requests/leave/{req['id']}/approve"
    assert (await client.put(url, headers=emp, json={})).status_code == 403
    pending = (await client.get("/api/admin/requests?type=leave&status=pending", headers=H["manager"])).json()
    assert pending["total"] == 1 and pending["items"][0]["user"]["employee_id"] == "EMP001"
    r = await client.put(url, headers=H["manager"], json={"comment": "enjoy"})
    assert r.status_code == 200 and r.json()["status"] == "approved" and r.json()["approver"]["employee_id"] == "MGR001"
    assert (await client.put(url, headers=H["manager"], json={})).json()["error"]["code"] == "NOT_PENDING"
    bal = (await client.get("/api/leave/balance", headers=emp)).json()["leave"][0]
    assert (bal["used"], bal["pending"], bal["remaining"]) == (2.0, 0.0, 10.0)
    assert any(o["payload"]["type"] == "leave_decision" for o in notifications.OUTBOX)
    # an approved leave blocks an overlapping second request
    r = await client.post(
        "/api/leave/requests",
        headers=emp,
        json={"leave_type_id": str(lt.id), "from_date": "2026-10-27", "to_date": "2026-10-27", "reason": "x"},
    )
    assert r.status_code == 409 and r.json()["error"]["code"] == "OVERLAPPING_REQUEST"

    # export: everything shows up for admin, employees are refused
    q = "type=all&from=2026-10-01&to=2026-10-31&include_reasons=true"
    r = await client.get(f"/api/admin/export?{q}", headers=H["admin"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Attendance", "Leave", "Late", "Permission", "On Duty"]
    att = sheet_rows(r.content, "Attendance")
    assert len(att) == 1 and att[0]["Employee ID"] == "EMP001" and att[0]["Status"] == "late" and att[0]["Minutes Late"] == 30
    lv = sheet_rows(r.content, "Leave")
    assert len(lv) == 1 and lv[0]["Status"] == "approved" and lv[0]["Days"] == 2 and lv[0]["Reason"] == "vacation"
    assert lv[0]["Decision Comment"] == "enjoy" and lv[0]["Approver"]
    lt_rows = sheet_rows(r.content, "Late")
    assert len(lt_rows) == 1 and lt_rows[0]["Minutes Late"] == 30
    assert sheet_rows(r.content, "Permission") == []
    # a filter that excludes the employee yields header-only sheets
    r = await client.get("/api/admin/export?type=leave&from=2026-12-01&to=2026-12-31", headers=H["admin"])
    assert r.status_code == 200 and sheet_rows(r.content, "Leave") == []
    assert (await client.get(f"/api/admin/export?{q}", headers=emp)).status_code == 403
    # the manager's export is team scoped (employee reports to the manager) but the reasons flag needs export.sensitive
    assert (await client.get(f"/api/admin/export?{q}", headers=H["manager"])).status_code == 403
    r = await client.get("/api/admin/export?type=attendance&from=2026-10-01&to=2026-10-31", headers=H["manager"])
    assert r.status_code == 200 and len(sheet_rows(r.content, "Attendance")) == 1
    # audit trail covers the chain
    actions = set((await db.execute(select(AuditLog.action))).scalars().all())
    assert {"leave.apply", "request.approve", "export.run"} <= actions


async def test_permission_request_then_late_exemption_flow(client, H, clock):
    """An approved permission covering shift start..to_time turns a would-be late punch into on-time."""
    emp = H["employee"]
    r = await client.post(
        "/api/permission/requests",
        headers=emp,
        json={"date": "2026-10-22", "from_time": "09:30", "to_time": "10:30", "reason": "bank"},
    )
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    assert r.json()["minutes"] == 60
    r = await client.put(f"/api/admin/requests/permission/{pid}/approve", headers=H["manager"], json={})
    assert r.status_code == 200 and r.json()["status"] == "approved"
    clock["now"] = dt.datetime(2026, 10, 22, 5, 0, tzinfo=dt.UTC)  # Thu 10:30 IST (grace already used by permission)
    r = await client.post("/api/attendance/check-in", headers=emp, json={**NEAR, "fix_time": clock["now"].isoformat()})
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "present" and r.json()["minutes_late"] == 0
    assert (await client.get("/api/late?from=2026-10-01&to=2026-10-31", headers=emp)).json()["total"] == 0
    bal = (await client.get("/api/leave/balance", headers=emp)).json()["permission"]
    assert bal["used_count"] == 1 and bal["used_minutes"] == 60


# ---------------------------------------------------------------- announcement -> read -> celebrations
async def test_announcement_read_and_celebrations_flow(client, H):
    emp, hr = H["employee"], H["hr"]
    r = await client.post(
        "/api/admin/announcements",
        headers=hr,
        json={"title": "Diwali", "body": "Office closed", "priority": "important", "target_all": True},
    )
    assert r.status_code == 201, r.text
    aid = r.json()["id"]
    assert (await client.get("/api/announcements/unread-count", headers=emp)).json() == {"unread_count": 1}
    assert (await client.get("/api/home", headers=emp)).json()["unread_announcements"] == 1
    # reading the detail does not mark as read; the explicit call does and is idempotent
    assert (await client.get(f"/api/announcements/{aid}", headers=emp)).json()["is_read"] is False
    assert (await client.get("/api/announcements/unread-count", headers=emp)).json()["unread_count"] == 1
    assert (await client.post(f"/api/announcements/{aid}/read", headers=emp)).status_code == 204
    assert (await client.post(f"/api/announcements/{aid}/read", headers=emp)).status_code == 204
    assert (await client.get("/api/announcements/unread-count", headers=emp)).json()["unread_count"] == 0
    assert (await client.get("/api/announcements?unread=true", headers=emp)).json()["total"] == 0
    assert (await client.get("/api/home", headers=emp)).json()["unread_announcements"] == 0
    reads = (await client.get(f"/api/admin/announcements/{aid}/reads", headers=hr)).json()
    assert reads["total"] == 1 and reads["read_count"] == 1 and reads["target_count"] == 4
    assert reads["items"][0]["user"]["employee_id"] == "EMP001"
    assert (await client.get(f"/api/admin/announcements/{aid}/reads", headers=emp)).status_code == 403
    adm_view = (await client.get("/api/admin/announcements", headers=hr)).json()["items"][0]
    assert adm_view["read_count"] == 1 and adm_view["status"] == "published"
    # soft delete hides it from employees (404, not 403)
    assert (await client.delete(f"/api/admin/announcements/{aid}", headers=hr)).status_code == 204
    assert (await client.get(f"/api/announcements/{aid}", headers=emp)).status_code == 404
    assert (await client.post(f"/api/announcements/{aid}/read", headers=emp)).status_code == 404

    # celebrations are opt-in and carry no age / birth year
    today = real_utcnow().date()
    soon = today + dt.timedelta(days=2)
    dob = dt.date(1990, soon.month, min(soon.day, 28)) if soon.month != 2 or soon.day <= 28 else dt.date(1990, 2, 28)
    assert (await client.put("/api/me/profile", headers=H["manager"], json={"dob": dob.isoformat()})).status_code == 200
    assert (await client.get("/api/home/celebrations?days=7", headers=emp)).json()["items"] == []  # not opted in
    assert (await client.put("/api/me/profile", headers=H["manager"], json={"show_celebrations": True})).status_code == 200
    items = (await client.get("/api/home/celebrations?days=7", headers=emp)).json()["items"]
    assert len(items) == 1 and items[0]["type"] == "birthday" and items[0]["years"] is None
    assert items[0]["user"]["employee_id"] == "MGR001"
    assert "1990" not in str(items[0])
    home = (await client.get("/api/home", headers=emp)).json()
    assert [c["user"]["employee_id"] for c in home["celebrations"]] == ["MGR001"]
    assert (await client.put("/api/me/profile", headers=H["manager"], json={"show_celebrations": False})).status_code == 200
    assert (await client.get("/api/home/celebrations?days=7", headers=emp)).json()["items"] == []


# ---------------------------------------------------------------- food yes -> cut-off -> admin count
async def test_food_flow_toggle_cutoff_and_admin_count(client, H, clock):
    emp, hr = H["employee"], H["hr"]
    r = await client.put("/api/food/enrollment", headers=emp, json={"date": "2026-10-22", "need_food": True})
    assert r.status_code == 200 and r.json()["need_food"] is True and r.json()["can_change"] is True
    r = await client.put("/api/food/enrollment", headers=H["manager"], json={"date": "2026-10-22", "need_food": True})
    assert r.status_code == 200
    r = await client.put("/api/food/enrollment", headers=emp, json={"date": "2026-10-22", "need_food": False})
    assert r.json()["need_food"] is False  # upsert, not duplicate
    daily = (await client.get("/api/admin/food/daily?date=2026-10-22", headers=hr)).json()
    assert daily["count_yes"] == 1 and daily["count_no"] == 1 and daily["count_unanswered"] == 2
    assert [e["user"]["employee_id"] for e in daily["employees"]] == ["MGR001"]
    assert (await client.get("/api/admin/food/daily?date=2026-10-22", headers=emp)).status_code == 403
    # after the cut-off (Wed 18:00 IST = 12:30Z) the change is refused and the count is frozen
    clock["now"] = dt.datetime(2026, 10, 21, 12, 30, tzinfo=dt.UTC)
    r = await client.put("/api/food/enrollment", headers=emp, json={"date": "2026-10-22", "need_food": True})
    assert r.status_code == 409 and r.json()["error"]["code"] == "FOOD_CUTOFF_PASSED"
    mine = (await client.get("/api/food/my-enrollments?from=2026-10-22&to=2026-10-22", headers=emp)).json()["items"][0]
    assert mine["need_food"] is False and mine["can_change"] is False
    assert (await client.get("/api/admin/food/daily?date=2026-10-22", headers=hr)).json()["count_yes"] == 1
    r = await client.get("/api/admin/export?type=food&from=2026-10-22&to=2026-10-22", headers=hr)
    assert r.status_code == 200
    assert [x["Employee ID"] for x in sheet_rows(r.content, "Food") if x["Need Food"] in ("Yes", True)] == ["MGR001"]


# ---------------------------------------------------------------- forced password change flow
async def test_new_employee_forced_password_change_flow(client, H, seeded):
    adm = H["admin"]
    r = await client.post(
        "/api/admin/employees",
        headers=adm,
        json={
            "employee_id": "EMP900",
            "full_name": "New Joiner",
            "official_email": "new.joiner@corp-mail.com",
            "department_id": str(seeded["dept"]),
            "shift_id": str(seeded["shift"]),
            "role_id": str(seeded["roles"]["employee"]),
            "join_date": "2026-10-01",
        },
    )
    assert r.status_code == 201, r.text
    temp = r.json()["temporary_password"]
    r = await do_login(client, "EMP900", temp)
    assert r.status_code == 200 and r.json()["must_change_password"] is True
    h = bearer(r.json())
    assert (await client.get("/api/auth/me", headers=h)).status_code == 200
    for path in ("/api/home", "/api/leave/types", "/api/attendance/today", "/api/announcements"):
        resp = await client.get(path, headers=h)
        assert resp.status_code == 403 and resp.json()["error"]["code"] == "PASSWORD_CHANGE_REQUIRED", path
    new_pw = "An0ther!Strong#Pass9"
    r = await client.post("/api/auth/change-password", headers=h, json={"current_password": temp, "new_password": new_pw})
    assert r.status_code == 200 and r.json()["must_change_password"] is False
    h2 = bearer(r.json())
    assert (await client.get("/api/home", headers=h2)).status_code == 200
    assert (await client.get("/api/home", headers=h)).status_code == 401  # old access token invalidated
    assert (await do_login(client, "EMP900", temp)).status_code == 401
    assert (await do_login(client, "EMP900", new_pw)).status_code == 200
